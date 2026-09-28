#!/usr/bin/env python
"""
Flat, filterable decision-maker export.

INPUT:  a consolidated shortlister report (``report_consolidated_*.xlsx``) whose
        "Decision Maker Contacts" sheet repeats the company on every contact row.
OUTPUT: an Excel workbook with AutoFilter on Company Name, so the built-in
        dropdown can slice the list per company.

Usage::

    python scripts/pivot_export.py --input var/runs/shortlister/report_consolidated_2026-08-05_23-12-56.xlsx
    python scripts/pivot_export.py --input report.xlsx --output pivot.xlsx --sheet "Decision Maker Contacts"
"""

import argparse
import os

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Columns to keep in the pivot (change order here if needed)
PIVOT_COLS = [
    "Company Name",
    "Decision Maker Name",
    "Decision Maker Position",
    "Decision Maker LinkedIn",
    "Location",
]

FONT = "Arial"
COLUMN_WIDTHS = {"A": 35, "B": 30, "C": 55, "D": 45, "E": 12}


def build_pivot(input_file: str, output_file: str, sheet_name: str) -> None:
    # ── READ ────────────────────────────────────────────────
    df = pd.read_excel(input_file, sheet_name=sheet_name)
    df = df[PIVOT_COLS].fillna("—")
    df = df.sort_values("Company Name").reset_index(drop=True)

    # ── STYLES ──────────────────────────────────────────────
    header_font = Font(name=FONT, bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    data_font = Font(name=FONT, size=10)
    border = Border(
        left=Side("thin", color="D9D9D9"), right=Side("thin", color="D9D9D9"),
        top=Side("thin", color="D9D9D9"), bottom=Side("thin", color="D9D9D9"),
    )
    alt_fill = PatternFill("solid", fgColor="F2F2F2")  # alternating row color
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(vertical="center", wrap_text=True)

    # ── BUILD WORKBOOK ──────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pivot - Filter by Company"

    # Header row
    for ci, col_name in enumerate(PIVOT_COLS, 1):
        cell = ws.cell(row=1, column=ci, value=col_name)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border

    # Data rows
    for ri, (_, row_data) in enumerate(df.iterrows(), 2):
        for ci, col_name in enumerate(PIVOT_COLS, 1):
            cell = ws.cell(row=ri, column=ci, value=row_data[col_name])
            cell.font, cell.border = data_font, border
            cell.alignment = center if ci == 1 else wrap
            if ri % 2 == 0:
                cell.fill = alt_fill

    # Column widths
    for col, w in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = w

    # ── AUTOFILTER (the actual filter dropdown) ─────────────
    last_row = ws.max_row
    last_col = get_column_letter(len(PIVOT_COLS))
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    # Freeze header row → stays visible when scrolling
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    # ── SAVE ────────────────────────────────────────────────
    wb.save(output_file)
    print(f"✓ Saved → {output_file}")
    print(f"  {len(df)} contacts | {df['Company Name'].nunique()} companies")
    print("  Open in Excel → click Company Name dropdown → pick company")


def main():
    parser = argparse.ArgumentParser(description="Build a filterable decision-maker pivot from a consolidated report")
    parser.add_argument("--input", required=True, help="Path to the consolidated report .xlsx")
    parser.add_argument(
        "--output",
        default=None,
        help="Destination .xlsx (default: <input>_pivot.xlsx next to the input file)",
    )
    parser.add_argument(
        "--sheet",
        default="Decision Maker Contacts",
        help='Sheet to read (default: "Decision Maker Contacts")',
    )
    args = parser.parse_args()

    if not os.path.exists(args.input):
        parser.error(f"input file not found: {args.input}")

    output_file = args.output
    if not output_file:
        stem, _ = os.path.splitext(args.input)
        output_file = f"{stem}_pivot.xlsx"

    build_pivot(args.input, output_file, args.sheet)


if __name__ == "__main__":
    main()
