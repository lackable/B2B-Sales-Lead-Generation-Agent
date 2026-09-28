"""Hermetic unit tests for the LinkedIn Finder LangGraph nodes.

No network, no LLM and no MCP: every module-level collaborator is monkeypatched.
"""

import json
import logging
from pathlib import Path

import httpx
import openpyxl
import pytest
from langchain_core.messages import HumanMessage, SystemMessage
from langgraph.checkpoint.memory import MemorySaver
from langgraph.graph import StateGraph
from langgraph.types import Command

from leadgen import config as app_config
from leadgen.agents.linkedin_finder import graph as linkedin_graph
from leadgen.agents.linkedin_finder.configuration import Configuration
from leadgen.agents.linkedin_finder.schemas import (
    DecisionMaker,
    DecisionMakerList,
    ResearchLoopDecision,
)
from leadgen.core.clients import bright_data as bright_data_module
from leadgen.core.telemetry import registry

SCRAPE_TOOL_NAME = "scrape_as_markdown"


# ──────────────────────────────────────────────────────────────────────────────
# Test doubles
# ──────────────────────────────────────────────────────────────────────────────
class EventRecorder:
    """Stand-in for the structured pipeline logger registered in the registry."""

    def __init__(self):
        self.records = []

    def event(self, event_type, data, *, level=None, component=None):
        self.records.append({"event_type": event_type, "data": data, "level": level, "component": component})
        return True

    def types(self):
        return [record["event_type"] for record in self.records]

    def structured_types(self):
        """The emitted node/tool events, without the ``log`` events from ``log_print``."""
        return [event_type for event_type in self.types() if event_type != "log"]

    def payloads(self, event_type):
        return [record["data"] for record in self.records if record["event_type"] == event_type]

    def single(self, event_type):
        payloads = self.payloads(event_type)
        assert len(payloads) == 1, f"expected exactly one {event_type!r} event, got {len(payloads)}"
        return payloads[0]


@pytest.fixture
def events(monkeypatch):
    recorder = EventRecorder()
    monkeypatch.setattr(registry, "_active_logger", recorder)
    return recorder


class FakeTool:
    def __init__(self, name, result="tool result", error=None):
        self.name = name
        self.result = result
        self.error = error
        self.calls = []

    async def ainvoke(self, args, config=None):
        self.calls.append(args)
        if self.error is not None:
            raise self.error
        return self.result


class FakeAIMessage:
    def __init__(self, tool_calls=None, content=""):
        self.tool_calls = tool_calls
        self.content = content


class StructuredStub:
    def __init__(self, schema):
        self.schema = schema
        self.retry_kwargs = None

    def with_retry(self, **kwargs):
        self.retry_kwargs = kwargs
        return self


class FakeChatModel:
    """Minimal stand-in for a ChatOpenAI instance."""

    def __init__(self):
        self.bound_tool_lists = []
        self.structured_stubs = []
        self.bound_model = FakeAIMessage(tool_calls=None)

    def bind_tools(self, tools):
        self.bound_tool_lists.append(tools)
        return self.bound_model

    def with_structured_output(self, schema):
        stub = StructuredStub(schema)
        self.structured_stubs.append(stub)
        return stub


class ScriptedInvoker:
    """Replaces ``invoke_model_with_rate_limit_retry`` with a deterministic queue."""

    def __init__(self, *responses):
        self.responses = list(responses)
        self.calls = []

    async def __call__(self, model_or_factory, input_data, **kwargs):
        self.calls.append({"model": model_or_factory, "messages": input_data, "kwargs": kwargs})
        outcome = self.responses.pop(0)
        if isinstance(outcome, BaseException):
            raise outcome
        return outcome


class FakeResponse:
    def __init__(self, payload=None, raise_for_status_error=None):
        self.payload = payload
        self.raise_for_status_error = raise_for_status_error

    def raise_for_status(self):
        if self.raise_for_status_error is not None:
            raise self.raise_for_status_error

    def json(self):
        return self.payload


def install_fake_httpx_client(monkeypatch, response=None, get_error=None, requests=None):
    """Replace ``httpx.AsyncClient`` with an async-context-manager double."""

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs):
            self.args = args
            self.kwargs = kwargs

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def get(self, url, params=None):
            if requests is not None:
                requests.append({"url": url, "params": params})
            if get_error is not None:
                raise get_error
            return response

    monkeypatch.setattr(linkedin_graph.httpx, "AsyncClient", FakeAsyncClient)


# ──────────────────────────────────────────────────────────────────────────────
# _parse_decision_makers_from_text
# ──────────────────────────────────────────────────────────────────────────────
class TestParseDecisionMakersFromText:
    def test_parses_fenced_json_list(self):
        text = (
            "Some preamble\n"
            '```json\n[{"name": "Jane Doe", "position": "CEO", "linkedin_url": "https://linkedin.com/in/jane"}]\n```\n'
            "trailing prose"
        )

        assert linkedin_graph._parse_decision_makers_from_text(text) == [
            {"name": "Jane Doe", "position": "CEO", "linkedin_url": "https://linkedin.com/in/jane"}
        ]

    def test_parses_fenced_dict_with_decision_makers_wrapper(self):
        text = '```json\n{"decision_makers": [{"name": "A"}, {"name": "B"}]}\n```'

        assert linkedin_graph._parse_decision_makers_from_text(text) == [{"name": "A"}, {"name": "B"}]

    def test_parses_single_fenced_dict_without_wrapper(self):
        text = '```json\n{"name": "Solo", "position": "Founder"}\n```'

        assert linkedin_graph._parse_decision_makers_from_text(text) == [{"name": "Solo", "position": "Founder"}]

    def test_parses_every_fenced_block_in_order(self):
        text = (
            '```json\n[{"name": "First"}]\n```\n'
            "noise between blocks\n"
            '```json\n{"decision_makers": [{"name": "Second"}]}\n```\n'
            '```json\n{"name": "Third"}\n```'
        )

        assert linkedin_graph._parse_decision_makers_from_text(text) == [
            {"name": "First"},
            {"name": "Second"},
            {"name": "Third"},
        ]

    def test_invalid_json_block_is_ignored_but_valid_blocks_survive(self):
        text = '```json\n{not: valid, json,}\n```\n```json\n[{"name": "Valid"}]\n```'

        assert linkedin_graph._parse_decision_makers_from_text(text) == [{"name": "Valid"}]

    def test_text_without_any_fenced_block_returns_empty_list(self):
        assert linkedin_graph._parse_decision_makers_from_text("no json here, just prose") == []
        assert linkedin_graph._parse_decision_makers_from_text("") == []

    def test_unfenced_json_is_not_picked_up(self):
        assert linkedin_graph._parse_decision_makers_from_text('[{"name": "unfenced"}]') == []

    def test_non_list_json_scalars_are_ignored(self):
        text = "```json\n123\n```\n```json\n\"a string\"\n```\n```json\nnull\n```\n```json\ntrue\n```\n```json\n1.5\n```"

        assert linkedin_graph._parse_decision_makers_from_text(text) == []

    def test_dict_with_non_list_decision_makers_key_is_kept_as_is(self):
        text = '```json\n{"decision_makers": "not-a-list"}\n```'

        assert linkedin_graph._parse_decision_makers_from_text(text) == [{"decision_makers": "not-a-list"}]


# ──────────────────────────────────────────────────────────────────────────────
# llm_query_generator
# ──────────────────────────────────────────────────────────────────────────────
class TestLlmQueryGenerator:
    @pytest.mark.asyncio
    async def test_returns_stripped_query_and_emits_events(self, monkeypatch, events):
        model = FakeChatModel()
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: model)
        monkeypatch.setattr(linkedin_graph, "get_today_str", lambda: "2024-01-02")

        captured = {}

        # ``invoke_model_with_rate_limit_retry`` is replaced by a plain coroutine function, so the
        # response object only needs the ``query`` attribute the node reads.
        class QueryResponse:
            query = "   site:linkedin.com/in Acme Corp CEO  "

        async def fake_invoke(model_or_factory, input_data, **kwargs):
            captured["built_model"] = model_or_factory("gpt-test-deployment")
            captured["messages"] = input_data
            return QueryResponse()

        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", fake_invoke)

        state = {
            "company_name": "Acme Corp",
            "company_linkedin": "https://www.linkedin.com/company/acme",
            "company_website": "https://acme.example",
            "location": "Mumbai",
        }
        result = await linkedin_graph.llm_query_generator(state, {"configurable": {"max_structured_output_retries": 7}})

        assert result == {"generated_query": "site:linkedin.com/in Acme Corp CEO"}
        assert captured["built_model"] is model.structured_stubs[-1]
        assert [stub.schema for stub in model.structured_stubs] == [linkedin_graph.GeneratedQuery]
        assert model.structured_stubs[-1].retry_kwargs == {"stop_after_attempt": 7}

        messages = captured["messages"]
        assert len(messages) == 1
        assert isinstance(messages[0], HumanMessage)
        prompt = messages[0].content
        assert "Company Name: Acme Corp" in prompt
        assert "Company LinkedIn: https://www.linkedin.com/company/acme" in prompt
        assert "Company Website: https://acme.example" in prompt
        assert "Location: Mumbai" in prompt
        assert "Today's Date: 2024-01-02" in prompt

        assert events.types()[:1] == ["node_enter"]
        assert events.single("search_query") == {
            "query_str": "site:linkedin.com/in Acme Corp CEO",
            "source": "serp_query_generator",
            "company": "Acme Corp",
        }
        node_enter = events.single("node_enter")
        assert node_enter["node_name"] == "llm_query_generator"
        assert node_enter["step_num"] == 1
        assert node_enter["company"] == "Acme Corp"
        assert node_enter["location"] == "Mumbai"
        node_exit = events.single("node_exit")
        assert node_exit["node_name"] == "llm_query_generator"
        assert node_exit["step_num"] == 1
        assert node_exit["output_summary"] == {"query": "site:linkedin.com/in Acme Corp CEO"}
        assert isinstance(node_exit["duration_ms"], int)
        assert events.structured_types() == ["node_enter", "search_query", "node_exit"]

    @pytest.mark.asyncio
    async def test_uses_retry_setting_from_runnable_config(self, monkeypatch):
        model = FakeChatModel()
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: model)

        class QueryResponse:
            query = "query text"

        async def fake_invoke(model_or_factory, input_data, **kwargs):
            model_or_factory("deployment")
            return QueryResponse()

        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", fake_invoke)

        await linkedin_graph.llm_query_generator({}, {"configurable": {"max_structured_output_retries": 9}})

        assert model.structured_stubs[-1].retry_kwargs == {"stop_after_attempt": 9}

    @pytest.mark.asyncio
    async def test_defaults_retries_when_config_missing(self, monkeypatch):
        model = FakeChatModel()
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: model)

        class QueryResponse:
            query = "query text"

        async def fake_invoke(model_or_factory, input_data, **kwargs):
            model_or_factory("deployment")
            return QueryResponse()

        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", fake_invoke)

        await linkedin_graph.llm_query_generator({}, {})

        assert model.structured_stubs[-1].retry_kwargs == {
            "stop_after_attempt": Configuration().max_structured_output_retries
        }

    @pytest.mark.asyncio
    async def test_missing_state_fields_fall_back_to_placeholders(self, monkeypatch, events):
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        prompts = []

        class QueryResponse:
            query = "q"

        async def fake_invoke(model_or_factory, input_data, **kwargs):
            prompts.append(input_data[0].content)
            return QueryResponse()

        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", fake_invoke)

        result = await linkedin_graph.llm_query_generator({"company_name": None, "location": ""}, {})

        assert result == {"generated_query": "q"}
        assert "Company Name: Unknown Company" in prompts[0]
        assert "Company LinkedIn: Not provided" in prompts[0]
        assert "Company Website: Not provided" in prompts[0]
        assert "Location: India" in prompts[0]
        assert events.single("node_enter")["company"] == "Unknown Company"
        assert events.single("node_enter")["location"] == "India"


# ──────────────────────────────────────────────────────────────────────────────
# serp_search
# ──────────────────────────────────────────────────────────────────────────────
def serp_state(query="site:linkedin.com/in Acme Corp CEO"):
    return {"company_name": "Acme Corp", "generated_query": query}


class TestSerpSearch:
    @pytest.mark.asyncio
    async def test_without_api_key_skips_the_call(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", None)

        def unexpected_client(*args, **kwargs):
            raise AssertionError("httpx.AsyncClient must not be constructed without an API key")

        monkeypatch.setattr(linkedin_graph.httpx, "AsyncClient", unexpected_client)

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result == {"serp_raw_response": {}, "ai_overview": "SerpAPI API Key not configured."}

    @pytest.mark.asyncio
    async def test_empty_api_key_also_skips_the_call(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "")
        monkeypatch.setattr(linkedin_graph.httpx, "AsyncClient", lambda *a, **k: pytest.fail("no HTTP expected"))

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result == {"serp_raw_response": {}, "ai_overview": "SerpAPI API Key not configured."}

    @pytest.mark.asyncio
    async def test_success_sends_expected_params_and_returns_ai_overview_text(self, monkeypatch, events):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "test-serpapi-key")
        payload = {
            "ai_overview": {"text": "Acme Corp is led by Jane Doe (CEO)."},
            "organic_results": [{"title": "t", "snippet": "s"}],
        }
        requests = []
        install_fake_httpx_client(monkeypatch, response=FakeResponse(payload=payload), requests=requests)

        result = await linkedin_graph.serp_search(serp_state("acme ceo"), {})

        assert result == {"serp_raw_response": payload, "ai_overview": "Acme Corp is led by Jane Doe (CEO)."}
        assert requests == [
            {
                "url": "https://serpapi.com/search",
                "params": {"q": "acme ceo", "api_key": "test-serpapi-key", "engine": "google", "num": 10},
            }
        ]
        node_exit = events.single("node_exit")
        assert node_exit["node_name"] == "serp_search"
        assert node_exit["step_num"] == 2
        assert node_exit["output_summary"] == {
            "ai_overview_chars": len("Acme Corp is led by Jane Doe (CEO)."),
            "organic_results": 1,
        }

    @pytest.mark.asyncio
    async def test_ai_overview_dict_without_text_is_json_dumped(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        payload = {"ai_overview": {"html": "<div>overview</div>"}}
        install_fake_httpx_client(monkeypatch, response=FakeResponse(payload=payload))

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result["ai_overview"] == json.dumps({"html": "<div>overview</div>"})
        assert result["serp_raw_response"] == payload

    @pytest.mark.asyncio
    async def test_ai_overview_as_plain_string_is_stringified(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        payload = {"ai_overview": "plain overview text"}
        install_fake_httpx_client(monkeypatch, response=FakeResponse(payload=payload))

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result["ai_overview"] == "plain overview text"

    @pytest.mark.asyncio
    async def test_answer_box_answer_takes_precedence(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        payload = {"answer_box": {"answer": "the answer", "snippet": "the snippet"}}
        install_fake_httpx_client(monkeypatch, response=FakeResponse(payload=payload))

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result["ai_overview"] == "the answer"

    @pytest.mark.asyncio
    async def test_answer_box_falls_back_to_snippet(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        payload = {"answer_box": {"answer": "", "snippet": "the snippet"}}
        install_fake_httpx_client(monkeypatch, response=FakeResponse(payload=payload))

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result["ai_overview"] == "the snippet"

    @pytest.mark.asyncio
    async def test_organic_results_become_bulleted_snippet_lines(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        payload = {
            "organic_results": [
                {"title": "T1", "snippet": "S1"},
                {"title": "no snippet"},
                {"title": "T2", "snippet": "S2"},
                {"title": "T3", "snippet": ""},
                {"title": "T4", "snippet": "S4"},
                {"title": "T5", "snippet": "S5"},
                {"title": "T6", "snippet": "S6"},
                {"title": "T7", "snippet": "S7"},
            ]
        }
        install_fake_httpx_client(monkeypatch, response=FakeResponse(payload=payload))

        result = await linkedin_graph.serp_search(serp_state(), {})

        # Only the first 5 organic results are inspected, and only ones carrying a snippet.
        assert result["ai_overview"] == "- T1: S1\n- T2: S2\n- T4: S4"
        assert result["serp_raw_response"] == payload

    @pytest.mark.asyncio
    async def test_empty_response_yields_empty_overview(self, monkeypatch):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        install_fake_httpx_client(monkeypatch, response=FakeResponse(payload={}))

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result == {"serp_raw_response": {}, "ai_overview": ""}

    @pytest.mark.asyncio
    async def test_http_error_returns_error_overview_and_empty_raw_response(self, monkeypatch, events):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        request = httpx.Request("GET", "https://serpapi.com/search")
        error = httpx.HTTPStatusError("502 Bad Gateway", request=request, response=httpx.Response(502, request=request))
        install_fake_httpx_client(monkeypatch, response=FakeResponse(raise_for_status_error=error))

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result["serp_raw_response"] == {}
        assert result["ai_overview"] == f"SerpAPI Error: {error}"
        assert result["ai_overview"].startswith("SerpAPI Error: ")
        tool_error = events.single("tool_error")
        assert tool_error == {"tool_name": "serp_search", "company": "Acme Corp", "error": str(error)}
        tool_error_record = next(r for r in events.records if r["event_type"] == "tool_error")
        assert tool_error_record["level"] == logging.WARNING
        assert events.payloads("node_exit") == []

    @pytest.mark.asyncio
    async def test_transport_exception_is_caught_too(self, monkeypatch, events):
        monkeypatch.setattr(app_config, "SERPAPI_API_KEY", "key")
        error = httpx.ConnectError("connection refused")
        install_fake_httpx_client(monkeypatch, get_error=error)

        result = await linkedin_graph.serp_search(serp_state(), {})

        assert result == {"serp_raw_response": {}, "ai_overview": f"SerpAPI Error: {error}"}
        assert events.single("tool_error")["error"] == str(error)


# ──────────────────────────────────────────────────────────────────────────────
# execute_tool_safely
# ──────────────────────────────────────────────────────────────────────────────
class TestExecuteToolSafely:
    @pytest.mark.asyncio
    async def test_success_returns_the_tool_result(self, monkeypatch):
        monkeypatch.setattr(bright_data_module, "reset_bright_data_tools", _unexpected_call("reset"))
        tool = FakeTool(SCRAPE_TOOL_NAME, result="scraped text")

        assert await linkedin_graph.execute_tool_safely(tool, {"url": "https://acme.example"}, {}) == "scraped text"
        assert tool.calls == [{"url": "https://acme.example"}]

    @pytest.mark.asyncio
    @pytest.mark.parametrize(
        "message",
        [
            "RemoteProtocolError: peer closed connection",
            "server disconnected",
            "502 Bad Gateway",
            "ReadTimeout: timed out",
            "httpcore.ReadError",
            "SSE post_writer closed",
        ],
    )
    async def test_connection_error_triggers_reconnect_and_retry(self, monkeypatch, message):
        resets = []

        async def fake_reset():
            resets.append(True)

        fresh_tool = FakeTool(SCRAPE_TOOL_NAME, result="fresh connection result")
        get_calls = []

        async def fake_get_tools(*args, **kwargs):
            get_calls.append(kwargs)
            return [FakeTool("other_tool"), fresh_tool]

        monkeypatch.setattr(bright_data_module, "reset_bright_data_tools", fake_reset)
        monkeypatch.setattr(bright_data_module, "get_bright_data_tools", fake_get_tools)

        broken_tool = FakeTool(SCRAPE_TOOL_NAME, error=RuntimeError(message))
        result = await linkedin_graph.execute_tool_safely(
            broken_tool, {"url": "https://acme.example"}, {}, company_name="Acme Corp"
        )

        assert result == "fresh connection result"
        assert len(resets) == 1
        assert get_calls == [{"force_refresh": True, "initial_backoff": 0.1}]
        assert broken_tool.calls == [{"url": "https://acme.example"}]
        assert fresh_tool.calls == [{"url": "https://acme.example"}]

    @pytest.mark.asyncio
    async def test_non_connection_error_returns_error_string_without_retry(self, monkeypatch):
        monkeypatch.setattr(bright_data_module, "reset_bright_data_tools", _unexpected_call("reset"))
        monkeypatch.setattr(bright_data_module, "get_bright_data_tools", _unexpected_call("get_bright_data_tools"))
        tool = FakeTool(SCRAPE_TOOL_NAME, error=ValueError("bad arguments"))

        result = await linkedin_graph.execute_tool_safely(tool, {"bad": True}, {})

        assert result == "Error executing tool: bad arguments"
        assert len(tool.calls) == 1

    @pytest.mark.asyncio
    async def test_connection_error_reports_error_when_no_fresh_tool_matches(self, monkeypatch):
        resets = []

        async def fake_reset():
            resets.append(True)

        get_calls = []

        async def fake_get_tools(*args, **kwargs):
            get_calls.append(kwargs)
            return [FakeTool("totally_different")]

        monkeypatch.setattr(bright_data_module, "reset_bright_data_tools", fake_reset)
        monkeypatch.setattr(bright_data_module, "get_bright_data_tools", fake_get_tools)
        tool = FakeTool(SCRAPE_TOOL_NAME, error=RuntimeError("server disconnected"))

        result = await linkedin_graph.execute_tool_safely(tool, {}, {})

        assert result == "Error executing tool: server disconnected"
        # Attempts 1 and 2 reconnect, attempt 3 exhausts max_retries and returns the error.
        assert len(resets) == 2
        assert get_calls == [{"force_refresh": True, "initial_backoff": 0.1}] * 2
        assert len(tool.calls) == 3

    @pytest.mark.asyncio
    async def test_fresh_tool_failure_falls_through_to_error_string(self, monkeypatch, capsys):
        resets = []

        async def fake_reset():
            resets.append(True)

        async def fake_get_tools(*args, **kwargs):
            return [FakeTool(SCRAPE_TOOL_NAME, error=RuntimeError("fresh connection also broken"))]

        monkeypatch.setattr(bright_data_module, "reset_bright_data_tools", fake_reset)
        monkeypatch.setattr(bright_data_module, "get_bright_data_tools", fake_get_tools)
        tool = FakeTool(SCRAPE_TOOL_NAME, error=RuntimeError("connection reset"))

        result = await linkedin_graph.execute_tool_safely(tool, {}, {}, max_retries=2)

        # The freshly fetched tool is retried, its failure is swallowed and the loop runs out of
        # attempts using the fresh (still broken) tool.
        assert result == "Error executing tool: fresh connection also broken"
        assert len(resets) == 1
        captured = capsys.readouterr().out
        assert "Instant reconnect attempt failed: fresh connection also broken" in captured
        assert "failed on attempt 2: fresh connection also broken" in captured

    @pytest.mark.asyncio
    async def test_tool_without_name_attribute_uses_str_repr(self, monkeypatch, capsys):
        class NamelessTool:
            async def ainvoke(self, args, config=None):
                raise ValueError("boom")

        monkeypatch.setattr(bright_data_module, "reset_bright_data_tools", _unexpected_call("reset"))

        result = await linkedin_graph.execute_tool_safely(NamelessTool(), {}, {})

        assert result == "Error executing tool: boom"
        assert "failed on attempt 1" in capsys.readouterr().out


def _unexpected_call(label):
    def _boom(*args, **kwargs):
        raise AssertionError(f"{label} must not be called")

    return _boom


# ──────────────────────────────────────────────────────────────────────────────
# research_loop
# ──────────────────────────────────────────────────────────────────────────────
def loop_state(**overrides):
    state = {
        "company_name": "Acme Corp",
        "company_website": "https://acme.example",
        "company_linkedin": "https://www.linkedin.com/company/acme",
        "location": "Mumbai",
        "ai_overview": "Acme is led by Jane Doe.",
        "decision_makers": [],
        "scrape_calls_used": 0,
        "loop_memory": [],
    }
    state.update(overrides)
    return state


def install_tools(monkeypatch, tools):
    async def fake_get_tools(*args, **kwargs):
        return tools

    monkeypatch.setattr(linkedin_graph, "get_bright_data_tools", fake_get_tools)


class TestResearchLoop:
    @pytest.mark.asyncio
    async def test_exits_when_scrape_budget_is_exhausted(self, monkeypatch, events):
        monkeypatch.setattr(app_config, "MAX_SCRAPE_CALLS", 3)
        monkeypatch.setattr(linkedin_graph, "get_bright_data_tools", _unexpected_call("get_bright_data_tools"))
        monkeypatch.setattr(linkedin_graph, "get_chat_model", _unexpected_call("get_chat_model"))

        cmd = await linkedin_graph.research_loop(loop_state(scrape_calls_used=3), {})

        assert isinstance(cmd, Command)
        assert cmd.goto == "output_node"
        assert cmd.update is None
        assert events.single("node_enter")["scrape_calls_used"] == 3
        node_exit = events.single("node_exit")
        assert node_exit["exit_reason"] == "budget_exhausted"
        assert node_exit["node_name"] == "research_loop"

    @pytest.mark.asyncio
    async def test_exceeds_scrape_budget_also_exits(self, monkeypatch):
        monkeypatch.setattr(app_config, "MAX_SCRAPE_CALLS", 3)
        monkeypatch.setattr(linkedin_graph, "get_bright_data_tools", _unexpected_call("get_bright_data_tools"))

        cmd = await linkedin_graph.research_loop(loop_state(scrape_calls_used=9), {})

        assert cmd.goto == "output_node"
        assert cmd.update is None

    @pytest.mark.asyncio
    async def test_exits_when_enough_valid_decision_makers_were_found(self, monkeypatch, events):
        monkeypatch.setattr(app_config, "MAX_DECISION_MAKERS", 2)
        monkeypatch.setattr(linkedin_graph, "get_bright_data_tools", _unexpected_call("get_bright_data_tools"))
        decision_makers = [
            {"name": "Jane", "position": "CEO", "linkedin_url": "https://linkedin.com/in/jane"},
            {"name": "", "position": "VP", "linkedin_url": "https://linkedin.com/in/noname"},
            {"name": "Company Page", "position": "N/A", "linkedin_url": "https://linkedin.com/company/acme"},
            {"name": "Joe", "position": "CTO", "linkedin_url": "https://LINKEDIN.com/IN/joe"},
        ]

        cmd = await linkedin_graph.research_loop(loop_state(decision_makers=decision_makers), {})

        assert cmd.goto == "output_node"
        assert cmd.update is None
        node_exit = events.single("node_exit")
        assert node_exit["exit_reason"] == "target_reached"
        assert events.single("node_enter")["decision_makers_count"] == 4

    @pytest.mark.asyncio
    async def test_tool_loading_failure_exits_to_output_node(self, monkeypatch, capsys):
        async def failing_get_tools(*args, **kwargs):
            raise RuntimeError("MCP down")

        monkeypatch.setattr(linkedin_graph, "get_bright_data_tools", failing_get_tools)

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        assert cmd.goto == "output_node"
        assert cmd.update is None
        assert "Could not load BrightData tools: MCP down" in capsys.readouterr().out

    @pytest.mark.asyncio
    async def test_blocks_scraping_linkedin_profile_urls(self, monkeypatch, events, capsys):
        tool = FakeTool(SCRAPE_TOOL_NAME, result="should never run")
        install_tools(monkeypatch, [tool])
        model = FakeChatModel()
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: model)
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"url": "https://www.linkedin.com/in/x"}}
        invoker = ScriptedInvoker(FakeAIMessage(tool_calls=[tool_call]))
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(scrape_calls_used=2), {})

        assert cmd.goto == "research_loop"
        assert cmd.update == {
            "scrape_calls_used": 3,
            "loop_memory": [
                {
                    "iteration": 3,
                    "tool_called": SCRAPE_TOOL_NAME,
                    "tool_input": {"url": "https://www.linkedin.com/in/x"},
                    "tool_output_preview": (
                        "BLOCKED: Scrape request to LinkedIn personal profile violated platform policy."
                    ),
                    "new_dms": [],
                    "reasoning": "Attempted forbidden LinkedIn profile scraping.",
                    "decision": "NEEDED",
                }
            ],
        }
        assert tool.calls == []
        assert len(invoker.calls) == 1
        assert "Blocked scraping LinkedIn profile URL: https://www.linkedin.com/in/x" in capsys.readouterr().out

    @pytest.mark.asyncio
    async def test_blocks_linkedin_url_passed_via_link_argument(self, monkeypatch):
        install_tools(monkeypatch, [FakeTool(SCRAPE_TOOL_NAME)])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"link": "HTTP://linkedin.com/in/y"}}
        monkeypatch.setattr(
            linkedin_graph, "invoke_model_with_rate_limit_retry", ScriptedInvoker(FakeAIMessage(tool_calls=[tool_call]))
        )

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        assert cmd.goto == "research_loop"
        assert cmd.update["loop_memory"][0]["tool_input"] == {"link": "HTTP://linkedin.com/in/y"}

    @pytest.mark.asyncio
    async def test_runs_tool_and_exits_with_not_needed_verdict(self, monkeypatch, events, capsys):
        tool_output = "Acme team page: Jane Doe, CEO, https://linkedin.com/in/janedoe"
        tool = FakeTool(SCRAPE_TOOL_NAME, result=tool_output)
        install_tools(monkeypatch, [tool])
        monkeypatch.setattr(linkedin_graph, "get_today_str", lambda: "2024-01-02")
        model = FakeChatModel()
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: model)
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"url": "https://acme.example/team"}}
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[tool_call]),
            DecisionMakerList(
                decision_makers=[
                    DecisionMaker(name="Jane Doe", position="CEO", linkedin_url="https://linkedin.com/in/janedoe")
                ]
            ),
            ResearchLoopDecision(status="NOT_NEEDED", reasoning="Enough decision makers gathered."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        new_dm = {"name": "Jane Doe", "position": "CEO", "linkedin_url": "https://linkedin.com/in/janedoe"}
        assert cmd.goto == "output_node"
        assert cmd.update == {
            "decision_makers": [new_dm],
            "scrape_calls_used": 1,
            "loop_memory": [
                {
                    "iteration": 1,
                    "tool_called": SCRAPE_TOOL_NAME,
                    "tool_input": {"url": "https://acme.example/team"},
                    "tool_output_preview": tool_output,
                    "new_dms": [new_dm],
                    "reasoning": "Enough decision makers gathered.",
                    "decision": "NOT_NEEDED",
                }
            ],
        }
        assert tool.calls == [{"url": "https://acme.example/team"}]
        assert model.bound_tool_lists == [[tool]]
        assert [stub.schema for stub in model.structured_stubs] == [DecisionMakerList, ResearchLoopDecision]
        assert invoker.calls[0]["model"] is model.bound_model
        assert invoker.calls[1]["model"].schema is DecisionMakerList
        assert invoker.calls[2]["model"].schema is ResearchLoopDecision

        system_message, human_message = invoker.calls[0]["messages"]
        assert isinstance(system_message, SystemMessage)
        assert isinstance(human_message, HumanMessage)
        system_prompt = system_message.content
        assert f"up to {app_config.MAX_DECISION_MAKERS} decision makers" in system_prompt
        assert "Acme Corp" in system_prompt
        assert "https://acme.example" in system_prompt
        assert "https://www.linkedin.com/company/acme" in system_prompt
        assert "Mumbai" in system_prompt
        assert "2024-01-02" in system_prompt
        assert "Acme is led by Jane Doe." in system_prompt
        assert f"- Tool Calls Remaining: {app_config.MAX_SCRAPE_CALLS}" in system_prompt
        assert "History of Previous Operations:\nNone" in system_prompt
        assert human_message.content == (
            "Select and call the best BrightData scraping tool to find decision makers or missing LinkedIn URLs."
        )

        extraction_prompt = invoker.calls[1]["messages"][0].content
        assert SCRAPE_TOOL_NAME in extraction_prompt
        assert "Acme Corp" in extraction_prompt
        assert "Mumbai" in extraction_prompt
        assert tool_output in extraction_prompt

        decision_prompt = invoker.calls[2]["messages"][0].content
        assert "Acme Corp" in decision_prompt
        assert f"- Tool Calls Used: 1 / {app_config.MAX_SCRAPE_CALLS}" in decision_prompt
        assert tool_output in decision_prompt

        assert events.single("tool_start") == {
            "tool_name": SCRAPE_TOOL_NAME,
            "company": "Acme Corp",
            "args_preview": "{'url': 'https://acme.example/team'}",
            "iteration": 1,
        }
        tool_end = events.single("tool_end")
        assert tool_end["tool_name"] == SCRAPE_TOOL_NAME
        assert tool_end["company"] == "Acme Corp"
        assert tool_end["output_chars"] == len(tool_output)
        assert isinstance(tool_end["latency_ms"], int)
        assert events.single("contact_found") == {
            "person_name": "Jane Doe",
            "role": "CEO",
            "linkedin_url": "https://linkedin.com/in/janedoe",
            "company": "Acme Corp",
            "source": SCRAPE_TOOL_NAME,
        }
        node_exit = events.single("node_exit")
        assert node_exit["exit_reason"] == "NOT_NEEDED"
        assert node_exit["new_dms"] == 1
        assert node_exit["total_dms"] == 1
        assert node_exit["iteration"] == 1
        assert "Exiting research_loop to output_node." in capsys.readouterr().out

    @pytest.mark.asyncio
    async def test_continues_loop_on_needed_verdict(self, monkeypatch, events):
        tool = FakeTool(SCRAPE_TOOL_NAME, result="page text")
        install_tools(monkeypatch, [tool])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"url": "https://acme.example/about"}}
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[tool_call]),
            DecisionMakerList(
                decision_makers=[
                    DecisionMaker(name="Joe", position="CTO", linkedin_url="https://linkedin.com/in/joe")
                ]
            ),
            ResearchLoopDecision(status="NEEDED", reasoning="Need more profiles."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(scrape_calls_used=4), {})

        assert cmd.goto == "research_loop"
        assert cmd.update["scrape_calls_used"] == 5
        assert cmd.update["decision_makers"] == [
            {"name": "Joe", "position": "CTO", "linkedin_url": "https://linkedin.com/in/joe"}
        ]
        assert cmd.update["loop_memory"][0]["iteration"] == 5
        assert cmd.update["loop_memory"][0]["decision"] == "NEEDED"
        assert events.single("node_exit")["exit_reason"] == "continuing"

    @pytest.mark.asyncio
    async def test_deduplicates_against_existing_decision_makers(self, monkeypatch, events):
        install_tools(monkeypatch, [FakeTool(SCRAPE_TOOL_NAME, result="page text")])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"url": "https://acme.example/team"}}
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[tool_call]),
            DecisionMakerList(
                decision_makers=[
                    DecisionMaker(name="Jane Doe", position="CEO", linkedin_url="https://LINKEDIN.com/in/janedoe"),
                    DecisionMaker(name="JANE DOE", position="CEO", linkedin_url="https://linkedin.com/in/other"),
                    DecisionMaker(name="Fresh Face", position="COO", linkedin_url="https://linkedin.com/in/fresh"),
                ]
            ),
            ResearchLoopDecision(status="NEEDED", reasoning="Keep going."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)
        existing = [{"name": "Jane Doe", "position": "CEO", "linkedin_url": "https://linkedin.com/in/janedoe"}]

        cmd = await linkedin_graph.research_loop(loop_state(decision_makers=existing), {})

        assert cmd.update["decision_makers"] == [
            {"name": "Fresh Face", "position": "COO", "linkedin_url": "https://linkedin.com/in/fresh"}
        ]
        assert len(events.payloads("contact_found")) == 1
        assert events.single("contact_found")["person_name"] == "Fresh Face"

    @pytest.mark.asyncio
    async def test_forced_exit_when_max_decision_makers_reached(self, monkeypatch):
        monkeypatch.setattr(app_config, "MAX_DECISION_MAKERS", 1)
        install_tools(monkeypatch, [FakeTool(SCRAPE_TOOL_NAME, result="page text")])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"url": "https://acme.example/team"}}
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[tool_call]),
            DecisionMakerList(
                decision_makers=[
                    DecisionMaker(name="Jane", position="CEO", linkedin_url="https://linkedin.com/in/jane")
                ]
            ),
            ResearchLoopDecision(status="NEEDED", reasoning="Still wants more."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        assert cmd.goto == "output_node"
        assert cmd.update["scrape_calls_used"] == 1

    @pytest.mark.asyncio
    async def test_forced_exit_when_scrape_budget_reached_after_the_call(self, monkeypatch):
        monkeypatch.setattr(app_config, "MAX_SCRAPE_CALLS", 1)
        install_tools(monkeypatch, [FakeTool(SCRAPE_TOOL_NAME, result="page text")])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"url": "https://acme.example/team"}}
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[tool_call]),
            DecisionMakerList(decision_makers=[]),
            ResearchLoopDecision(status="NEEDED", reasoning="Still wants more."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        assert cmd.goto == "output_node"
        assert cmd.update["scrape_calls_used"] == 1
        assert cmd.update["decision_makers"] == []

    @pytest.mark.asyncio
    async def test_missing_tool_reports_not_found_and_continues(self, monkeypatch, events):
        install_tools(monkeypatch, [FakeTool("search_engine", result="other")])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        tool_call = {"name": "missing_tool", "args": {"url": "https://acme.example"}}
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[tool_call]),
            DecisionMakerList(decision_makers=[]),
            ResearchLoopDecision(status="NEEDED", reasoning="Try something else."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        assert cmd.goto == "research_loop"
        assert cmd.update["loop_memory"][0]["tool_output_preview"] == "Error: Tool 'missing_tool' not found."
        assert invoker.calls[1]["messages"][0].content.count("Error: Tool 'missing_tool' not found.") == 1
        assert events.single("tool_error") == {
            "tool_name": "missing_tool",
            "error": "Error: Tool 'missing_tool' not found.",
            "company": "Acme Corp",
        }

    @pytest.mark.asyncio
    async def test_structured_extraction_failure_is_swallowed(self, monkeypatch, events, capsys):
        install_tools(monkeypatch, [FakeTool(SCRAPE_TOOL_NAME, result="page text")])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        tool_call = {"name": SCRAPE_TOOL_NAME, "args": {"url": "https://acme.example/team"}}
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[tool_call]),
            ValueError("structured output failed"),
            ResearchLoopDecision(status="NEEDED", reasoning="Extraction failed, retry."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        assert cmd.goto == "research_loop"
        assert cmd.update["decision_makers"] == []
        assert cmd.update["loop_memory"][0]["decision"] == "NEEDED"
        assert events.payloads("contact_found") == []
        assert "Failed structured extraction from tool output: structured output failed" in capsys.readouterr().out

    @pytest.mark.asyncio
    async def test_no_tool_call_with_not_needed_decision_exits(self, monkeypatch, events):
        install_tools(monkeypatch, [FakeTool(SCRAPE_TOOL_NAME)])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=[]),
            ResearchLoopDecision(status="NOT_NEEDED", reasoning="Nothing else to do."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(scrape_calls_used=2), {})

        assert cmd.goto == "output_node"
        assert cmd.update is None
        decision_prompt = invoker.calls[1]["messages"][0].content
        assert "No tool executed in last step." in decision_prompt
        assert f"- Tool Calls Used: 2 / {app_config.MAX_SCRAPE_CALLS}" in decision_prompt
        assert invoker.calls[1]["model"].schema is ResearchLoopDecision
        # The no-tool-call shortcut returns without emitting a node_exit event.
        assert events.payloads("node_exit") == []

    @pytest.mark.asyncio
    async def test_no_tool_call_with_needed_decision_loops(self, monkeypatch, events):
        install_tools(monkeypatch, [FakeTool(SCRAPE_TOOL_NAME)])
        monkeypatch.setattr(linkedin_graph, "get_chat_model", lambda *args, **kwargs: FakeChatModel())
        invoker = ScriptedInvoker(
            FakeAIMessage(tool_calls=None),
            ResearchLoopDecision(status="NEEDED", reasoning="Keep researching."),
        )
        monkeypatch.setattr(linkedin_graph, "invoke_model_with_rate_limit_retry", invoker)

        cmd = await linkedin_graph.research_loop(loop_state(), {})

        assert cmd.goto == "research_loop"
        assert cmd.update is None
        assert events.payloads("node_exit") == []
        decision_prompt = invoker.calls[1]["messages"][0].content
        assert "No tool executed in last step." in decision_prompt
        assert invoker.calls[1]["model"].schema is ResearchLoopDecision


# ──────────────────────────────────────────────────────────────────────────────
# output_node
# ──────────────────────────────────────────────────────────────────────────────
def output_state(**overrides):
    state = {
        "company_name": "Acme Corp",
        "company_website": "https://acme.example",
        "company_linkedin": "https://www.linkedin.com/company/acme",
        "location": "Mumbai",
        "session_id": "sess1",
        "generated_query": "site:linkedin.com/in Acme Corp CEO",
        "ai_overview": "Overview text",
        "serp_raw_response": {"organic_results": [{"title": "t"}]},
        "scrape_calls_used": 2,
        "decision_makers": [
            {"name": "Jane Doe", "position": "CEO", "linkedin_url": "https://linkedin.com/in/janedoe"},
            {"name": "Jane Duplicate", "position": "CEO", "linkedin_url": "https://linkedin.com/in/janedoe"},
            {"name": "Joe Roe", "position": "CTO", "linkedin_url": "https://linkedin.com/in/joeroe"},
            {"name": "No Link", "position": "VP", "linkedin_url": ""},
            {"name": "Blank Link", "position": "VP", "linkedin_url": "   "},
        ],
        "loop_memory": [{"iteration": 1, "tool_called": SCRAPE_TOOL_NAME}],
    }
    state.update(overrides)
    return state


class TestOutputNode:
    @pytest.mark.asyncio
    async def test_writes_excel_and_json_and_emits_paths(self, isolated_var, events):
        result = await linkedin_graph.output_node(output_state(), {})

        assert result == {}
        output_dir = Path(app_config.OUTPUT_DIR)
        excel_files = sorted(output_dir.glob("contacts_*.xlsx"))
        assert [f.name for f in excel_files] == ["contacts_Acme_Corp_sess1.xlsx"]
        excel_path = excel_files[0]
        json_path = output_dir / "agent_state_sess1.json"
        assert json_path.exists()

        expected_dms = [
            {"name": "Jane Doe", "position": "CEO", "linkedin_url": "https://linkedin.com/in/janedoe"},
            {"name": "Joe Roe", "position": "CTO", "linkedin_url": "https://linkedin.com/in/joeroe"},
        ]
        payload = json.loads(json_path.read_text(encoding="utf-8"))
        assert payload == {
            "session_id": "sess1",
            "company_name": "Acme Corp",
            "company_website": "https://acme.example",
            "company_linkedin": "https://www.linkedin.com/company/acme",
            "location": "Mumbai",
            "generated_query": "site:linkedin.com/in Acme Corp CEO",
            "ai_overview": "Overview text",
            "serp_raw_response": {"organic_results": [{"title": "t"}]},
            "scrape_calls_used": 2,
            "decision_makers": expected_dms,
            "loop_memory": [{"iteration": 1, "tool_called": SCRAPE_TOOL_NAME}],
        }

        workbook = openpyxl.load_workbook(str(excel_path))
        sheet = workbook.active
        assert sheet.title == "Decision Makers"
        assert [cell.value for cell in sheet[1]] == [
            "Company Name",
            "Company Website",
            "Company LinkedIn",
            "Location",
            "Decision Maker Position",
            "Decision Maker Name",
            "Decision Maker LinkedIn",
        ]
        assert sheet.max_row == 3
        assert sheet.cell(row=2, column=6).value == "Jane Doe"
        assert sheet.cell(row=3, column=6).value == "Joe Roe"

        node_enter = events.single("node_enter")
        assert node_enter["node_name"] == "output_node"
        assert node_enter["company"] == "Acme Corp"
        node_exit = events.single("node_exit")
        assert node_exit["output_summary"] == {
            "decision_makers_count": 2,
            "excel_path": str(excel_path),
            "json_path": str(json_path),
        }
        assert node_exit["node_name"] == "output_node"
        assert node_exit["step_num"] == 4

    @pytest.mark.asyncio
    async def test_default_session_id_is_used_when_missing(self, isolated_var):
        state = output_state()
        state.pop("session_id")
        await linkedin_graph.output_node(state, {})

        output_dir = Path(app_config.OUTPUT_DIR)
        assert (output_dir / "agent_state_default_session.json").exists() is True
        assert (output_dir / "contacts_Acme_Corp_default_session.xlsx").exists() is True

    @pytest.mark.asyncio
    async def test_json_write_failure_is_swallowed(self, isolated_var, monkeypatch, capsys, events):
        def exploding_dump(*args, **kwargs):
            raise OSError("disk full")

        monkeypatch.setattr(linkedin_graph.json, "dump", exploding_dump)

        result = await linkedin_graph.output_node(output_state(), {})

        assert result == {}
        output_dir = Path(app_config.OUTPUT_DIR)
        assert (output_dir / "contacts_Acme_Corp_sess1.xlsx").exists() is True
        assert "Failed to write JSON agent state: disk full" in capsys.readouterr().out
        assert events.single("node_exit")["output_summary"]["json_path"] == str(output_dir / "agent_state_sess1.json")


# ──────────────────────────────────────────────────────────────────────────────
# graph construction
# ──────────────────────────────────────────────────────────────────────────────
class TestGraphConstruction:
    def test_build_graph_with_explicit_checkpointer_returns_compiled_graph(self):
        saver = MemorySaver()

        compiled = linkedin_graph.build_graph(saver)

        assert compiled is not linkedin_graph.contact_finder
        assert hasattr(compiled, "ainvoke")
        node_names = set(compiled.get_graph().nodes)
        assert {"llm_query_generator", "serp_search", "research_loop", "output_node"} <= node_names

    def test_build_graph_defaults_to_the_linkedin_checkpointer_path(self, monkeypatch):
        recorded = []
        saver = MemorySaver()

        def fake_get_checkpointer(path):
            recorded.append(path)
            return saver

        monkeypatch.setattr(linkedin_graph, "get_checkpointer", fake_get_checkpointer)

        compiled = linkedin_graph.build_graph()

        assert recorded == [app_config.LINKEDIN_MEMORY_DB_PATH]
        assert hasattr(compiled, "ainvoke")

    def test_module_level_builder_and_compiled_graph_exist(self):
        assert isinstance(linkedin_graph.contact_finder_builder, StateGraph)
        assert hasattr(linkedin_graph.contact_finder, "ainvoke")
        assert hasattr(linkedin_graph.contact_finder, "get_graph")
        node_names = set(linkedin_graph.contact_finder.get_graph().nodes)
        assert {"llm_query_generator", "serp_search", "research_loop", "output_node"} <= node_names
