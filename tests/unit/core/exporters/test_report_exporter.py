import json
import os

from openpyxl import load_workbook

from leadgen import config
from leadgen.core.exporters.report import export_consolidated_excel, export_excel, save_all


def test_export_excel_verified_and_non_verified_status(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_OUTPUT_DIR", str(tmp_path))
    sample_companies = [
        {
            "name": "Verified Tech Ltd",
            "ticker": "VT",
            "industry": "Software",
            "revenue_ttm": 1000000,
            "net_income_ttm": 200000,
            "employee_count": 500,
            "business_summary": "Tech company",
            "website": "https://verifiedtech.com",
            "linkedin_guessed": "https://linkedin.com/company/verified-tech",
            "linkedin_verified": "https://linkedin.com/company/verified-tech",
            "linkedin_confirmed": True
        },
        {
            "name": "Fallback Tech Ltd",
            "ticker": "FT",
            "industry": "Software",
            "revenue_ttm": 500000,
            "net_income_ttm": 100000,
            "employee_count": 200,
            "business_summary": "Fallback company",
            "website": "https://fallbacktech.com",
            "linkedin_guessed": "https://linkedin.com/company/fallback-tech",
            "linkedin_verified": "https://linkedin.com/company/fallback-tech",
            "linkedin_confirmed": False  # Fallback case
        }
    ]
    
    filepath = export_excel(sample_companies, timestamp="test_export")
    assert os.path.exists(filepath)
    
    wb = load_workbook(filepath)
    ws = wb.active
    
    # Header row is row 1. Row 2 is Verified Tech, Row 3 is Fallback Tech.
    # Column 11 is 'LinkedIn Confirmed?' (column 10 holds the LinkedIn hyperlink)
    row2_status = ws.cell(row=2, column=11).value
    row3_status = ws.cell(row=3, column=11).value
    
    assert row2_status == "Verified"
    assert row3_status == "Non-Verified"


# ── export_consolidated_excel ─────────────────────────────────────────────────

def test_export_consolidated_excel_two_sheets(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_OUTPUT_DIR", str(tmp_path))
    companies = [{"name": "Acme", "linkedin_confirmed": True}]

    path = export_consolidated_excel(companies, timestamp="c1")

    assert os.path.exists(path)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Verified Shortlist", "Decision Maker Contacts"]


def test_export_consolidated_excel_verified_and_non_verified_styling(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_OUTPUT_DIR", str(tmp_path))
    companies = [
        {
            "name": "Verified Ltd",
            "linkedin_confirmed": True,
            "linkedin_verified": "https://linkedin.com/company/verified",
        },
        {"name": "Unverified Ltd", "linkedin_confirmed": False},
    ]

    path = export_consolidated_excel(companies, timestamp="style")
    ws1 = load_workbook(path)["Verified Shortlist"]

    verified = ws1.cell(row=2, column=11)
    assert verified.value == "Verified"
    assert verified.font.color.rgb == "000F5132"
    assert verified.fill.start_color.rgb == "00D1E7DD"

    non_verified = ws1.cell(row=3, column=11)
    assert non_verified.value == "Non-Verified"
    assert non_verified.font.color.rgb == "00842029"
    assert non_verified.fill.start_color.rgb == "00F8D7DA"


def test_export_consolidated_excel_placeholder_contact_for_company_without_dms(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_OUTPUT_DIR", str(tmp_path))
    companies = [{"name": "NoDm Ltd", "location": "India", "linkedin_confirmed": False}]

    path = export_consolidated_excel(companies, timestamp="nodm")
    ws2 = load_workbook(path)["Decision Maker Contacts"]

    # row 1 = banner, row 2 = sub-header, row 3 = placeholder contact
    assert ws2.cell(row=1, column=1).value == "  ▸ NoDm Ltd  (1 contacts)"
    assert ws2.cell(row=3, column=2).value == "—"
    assert ws2.cell(row=3, column=3).value == "—"
    assert ws2.cell(row=3, column=4).value == "—"
    assert ws2.cell(row=3, column=6).value == "India"


def test_export_consolidated_excel_applies_row_grouping(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_OUTPUT_DIR", str(tmp_path))
    companies = [
        {
            "name": "Grp Ltd",
            "linkedin_confirmed": True,
            "decision_makers": [{"name": "A"}, {"name": "B"}],
        }
    ]

    path = export_consolidated_excel(companies, timestamp="grp")
    ws2 = load_workbook(path)["Decision Maker Contacts"]

    outlined = [rd for rd in ws2.row_dimensions.values() if getattr(rd, "outlineLevel", 0) >= 1]
    assert outlined, "expected grouped contact rows to carry an outline level"


def test_export_consolidated_excel_hyperlinks_contact_linkedin(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_OUTPUT_DIR", str(tmp_path))
    companies = [
        {
            "name": "Acme",
            "linkedin_confirmed": True,
            "linkedin_verified": "https://linkedin.com/company/acme",
            "decision_makers": [{"name": "Jane", "linkedin_url": "https://linkedin.com/in/jane"}],
        }
    ]

    path = export_consolidated_excel(companies, timestamp="links")
    ws2 = load_workbook(path)["Decision Maker Contacts"]

    assert ws2.cell(row=3, column=4).hyperlink is not None
    assert ws2.cell(row=3, column=4).hyperlink.target == "https://linkedin.com/in/jane"
    assert ws2.cell(row=3, column=5).hyperlink is not None
    assert ws2.cell(row=3, column=5).hyperlink.target == "https://linkedin.com/company/acme"


# ── save_all ──────────────────────────────────────────────────────────────────

def test_save_all_writes_files_and_returns_paths(isolated_var):
    companies = [
        {"name": "Acme", "linkedin_confirmed": True, "decision_makers": [{"name": "Jane"}]},
    ]

    result = save_all(companies, screener_query={"market": "india"}, session_id="sess-1")

    assert set(result) == {"timestamp", "excel", "shortlist_excel", "json"}
    assert os.path.exists(result["excel"])
    assert os.path.exists(result["shortlist_excel"])
    assert os.path.exists(result["json"])

    assert result["excel"].startswith(config.EXCEL_OUTPUT_DIR)
    assert result["shortlist_excel"].startswith(config.EXCEL_OUTPUT_DIR)
    assert result["json"].startswith(config.REPORTS_OUTPUT_DIR)

    assert os.path.basename(result["excel"]).startswith("report_consolidated_")
    assert os.path.basename(result["shortlist_excel"]).startswith("report_")


def test_save_all_json_snapshot_payload(isolated_var):
    companies = [{"name": "Acme", "linkedin_confirmed": False}]

    result = save_all(companies, screener_query={"market": "us"}, session_id="sess-2")

    with open(result["json"], encoding="utf-8") as handle:
        payload = json.load(handle)

    assert payload["session_id"] == "sess-2"
    assert payload["timestamp"] == result["timestamp"]
    assert payload["screener_query"] == {"market": "us"}
    assert payload["total_companies"] == 1
    assert payload["companies"] == companies


def test_save_all_defaults_screener_query_to_empty_dict(isolated_var):
    result = save_all([], session_id="sess-3")

    with open(result["json"], encoding="utf-8") as handle:
        payload = json.load(handle)

    assert payload["screener_query"] == {}
    assert payload["total_companies"] == 0


def test_save_all_returns_error_when_output_dir_unusable(isolated_var, tmp_path, monkeypatch):
    blocker = tmp_path / "blocker"
    blocker.write_text("not a directory", encoding="utf-8")
    # A directory cannot live under an existing *file*.
    monkeypatch.setattr(config, "EXCEL_OUTPUT_DIR", str(blocker / "sub"))

    result = save_all([], session_id="s")

    assert "error" in result
