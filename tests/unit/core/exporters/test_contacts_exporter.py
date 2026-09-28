"""Unit tests for :mod:`leadgen.core.exporters.contacts`."""

import os

from openpyxl import load_workbook

from leadgen import config
from leadgen.core.exporters.contacts import export_contacts_excel

HEADERS = [
    "Company Name",
    "Company Website",
    "Company LinkedIn",
    "Location",
    "Decision Maker Position",
    "Decision Maker Name",
    "Decision Maker LinkedIn",
]


def _row(ws, row):
    return [ws.cell(row=row, column=col).value for col in range(1, 8)]


# ── Filename construction ─────────────────────────────────────────────────────

def test_filename_uses_session_when_company_contained(isolated_var):
    path = export_contacts_excel("Acme Corp", "", "", "", [], "run_acme_corp_2024")
    assert os.path.basename(path) == "contacts_run_acme_corp_2024.xlsx"


def test_filename_includes_safe_company_when_not_contained(isolated_var):
    path = export_contacts_excel("Acme Corp", "", "", "", [], "session123")
    assert os.path.basename(path) == "contacts_Acme_Corp_session123.xlsx"


def test_filename_strips_unsafe_characters(isolated_var):
    path = export_contacts_excel("Foo/Bar:Baz*", "", "", "", [], "s1")
    assert os.path.basename(path) == "contacts_FooBarBaz_s1.xlsx"


def test_filename_uses_company_placeholder_for_empty_name(isolated_var):
    path = export_contacts_excel("", "", "", "", [], "s1")
    assert os.path.basename(path) == "contacts_Company_s1.xlsx"


def test_file_written_under_configured_output_dir(isolated_var):
    path = export_contacts_excel("Acme", "", "", "", [], "s1")
    assert os.path.dirname(path) == str(config.OUTPUT_DIR)
    assert os.path.exists(path)


# ── Sheet layout ──────────────────────────────────────────────────────────────

def test_sheet_title_and_headers(isolated_var):
    path = export_contacts_excel("Acme", "", "", "India", [], "sess")
    ws = load_workbook(path).active

    assert ws.title == "Decision Makers"
    assert _row(ws, 1) == HEADERS


def test_row_values_without_urls(isolated_var):
    dms = [{"position": "CEO", "name": "Jane Doe"}, {"position": "", "name": "John"}]
    path = export_contacts_excel("Acme", "", "", "India", dms, "sess")
    ws = load_workbook(path).active

    row2 = _row(ws, 2)
    assert row2[0] == "Acme"
    assert row2[1] in (None, "")  # empty website
    assert row2[2] in (None, "")  # empty company linkedin
    assert row2[3] == "India"
    assert row2[4] == "CEO"
    assert row2[5] == "Jane Doe"
    assert row2[6] in (None, "")  # empty DM linkedin

    row3 = _row(ws, 3)
    assert row3[4] == "Unknown"  # empty position falls back
    assert row3[5] == "John"


def test_default_position_for_missing_and_none(isolated_var):
    dms = [{"name": "A"}, {"name": "B", "position": None}]
    path = export_contacts_excel("Acme", "", "", "", dms, "s")
    ws = load_workbook(path).active
    assert ws.cell(row=2, column=5).value == "Unknown"
    assert ws.cell(row=3, column=5).value == "Unknown"


# ── Hyperlinks ────────────────────────────────────────────────────────────────

def test_hyperlink_formulas_written_when_urls_present(isolated_var):
    dms = [{"position": "CTO", "name": "Amy", "linkedin_url": "https://linkedin.com/in/amy"}]
    path = export_contacts_excel(
        "Acme",
        "https://acme.com",
        "https://linkedin.com/company/acme",
        "India",
        dms,
        "s",
    )
    ws = load_workbook(path).active

    assert ws.cell(row=2, column=2).value == '=HYPERLINK("https://acme.com", "https://acme.com")'
    assert ws.cell(row=2, column=3).value == (
        '=HYPERLINK("https://linkedin.com/company/acme", "https://linkedin.com/company/acme")'
    )
    assert ws.cell(row=2, column=7).value == '=HYPERLINK("https://linkedin.com/in/amy", "https://linkedin.com/in/amy")'

    # Link styling applied
    for col in (2, 3, 7):
        assert ws.cell(row=2, column=col).font.underline == "single"


def test_no_hyperlink_formulas_when_urls_absent(isolated_var):
    dms = [{"name": "Amy", "position": "CTO"}]
    path = export_contacts_excel("Acme", "", "", "", dms, "s")
    ws = load_workbook(path).active

    for col in (2, 3, 7):
        assert ws.cell(row=2, column=col).value in (None, "")


# ── Column widths ─────────────────────────────────────────────────────────────

def test_column_widths_at_least_15(isolated_var):
    dms = [{"name": "Amy", "position": "CEO"}]
    path = export_contacts_excel("Acme", "https://acme.com", "", "India", dms, "s")
    ws = load_workbook(path).active

    for col_letter in "ABCDEFG":
        assert ws.column_dimensions[col_letter].width >= 15
