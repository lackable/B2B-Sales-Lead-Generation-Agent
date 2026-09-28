"""Unit tests for ``leadgen.api.parsers`` — URL helpers, Hunter count, Excel parsing,
session consolidation lookup, decision-maker extraction and location extraction."""

import json
import os

import pytest
from openpyxl import Workbook

from leadgen import config
from leadgen.api import parsers
from leadgen.api.parsers import (
    extract_domain_from_url,
    extract_location_from_query,
    extract_url,
    get_all_decision_makers,
    get_current_session_consolidated_excel,
    hunter_email_count,
    parse_shortlister_excel,
)
from leadgen.api.state import shortlister_state

SHORTLISTER_HEADERS = [
    "Company Name", "Ticker", "Industry", "Revenue (TTM)", "Net Income (TTM)",
    "Employee Count", "Location", "Business Summary", "Website",
    "LinkedIn (Verified)", "LinkedIn Confirmed?",
]

FLAT_DM_HEADERS = [
    "Company Name", "Company Website", "Company LinkedIn", "Location",
    "Decision Maker Position", "Decision Maker Name", "Decision Maker LinkedIn",
]


# ── helpers ────────────────────────────────────────────────────────────────────

def _holes_row(**overrides):
    row = {
        "Company Name": "Acme Ltd", "Ticker": "ACME", "Industry": "Tech",
        "Revenue (TTM)": 1000, "Net Income (TTM)": 100, "Employee Count": 50,
        "Location": "Mumbai", "Business Summary": "Great company",
        "Website": "https://acmecorp.com", "LinkedIn (Verified)": "linkedin.com/company/acme",
        "LinkedIn Confirmed?": True,
    }
    row.update(overrides)
    return [row[h] for h in SHORTLISTER_HEADERS]


def _write_shortlister_workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    for ci, header in enumerate(SHORTLISTER_HEADERS, 1):
        ws.cell(row=1, column=ci, value=header)
    for ri, row in enumerate(rows, 2):
        for ci, value in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=value)
    wb.save(path)
    return path


def _write_consolidated_flat(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Decision Maker Contacts"
    for ci, header in enumerate(FLAT_DM_HEADERS, 1):
        ws.cell(row=1, column=ci, value=header)
    for ri, row in enumerate(rows, 2):
        for ci, value in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=value)
    wb.save(path)
    return path


def _touch(path, mtime=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"x")
    if mtime is not None:
        os.utime(path, (mtime, mtime))
    return path


# ── extract_url ────────────────────────────────────────────────────────────────

def test_extract_url_none_and_empty():
    assert extract_url(None) == ""
    assert extract_url("") == ""


def test_extract_url_cell_hyperlink_full_url_unchanged():
    ws = Workbook().active
    cell = ws.cell(row=1, column=1, value="Acme")
    cell.hyperlink = "https://acmecorp.com"
    assert extract_url(cell) == "https://acmecorp.com"


def test_extract_url_cell_hyperlink_bare_domain_prepends_scheme():
    ws = Workbook().active
    cell = ws.cell(row=1, column=1, value=None)
    cell.hyperlink = "linkedin.com/company/acme"
    assert extract_url(cell) == "https://linkedin.com/company/acme"


def test_extract_url_cell_hyperlink_unrelated_target_falls_through_to_value():
    ws = Workbook().active
    cell = ws.cell(row=1, column=1, value="unrelated text")
    cell.hyperlink = "acmecorp.com"
    assert extract_url(cell) == ""


def test_extract_url_plain_url_string():
    assert extract_url("https://acmecorp.com/team") == "https://acmecorp.com/team"


def test_extract_url_linkedin_without_scheme():
    assert extract_url("linkedin.com/in/john-doe") == "https://linkedin.com/in/john-doe"


def test_extract_url_www_value_gets_scheme():
    assert extract_url("www.acmecorp.com") == "https://www.acmecorp.com"


def test_extract_url_hyperlink_formula():
    assert extract_url('=HYPERLINK("https://acmecorp.com","Acme Corp")') == "https://acmecorp.com"


def test_extract_url_unrelated_text():
    assert extract_url("Just some notes here") == ""


@pytest.mark.parametrize("value", ["N/A", "not available", "not found", "none", "—", "", None])
def test_extract_url_placeholders(value):
    assert extract_url(value) == ""


# ── extract_domain_from_url ────────────────────────────────────────────────────

@pytest.mark.parametrize("raw, expected", [
    ("https://www.acmecorp.com/path?q=1", "acmecorp.com"),
    ("acmecorp.com", "acmecorp.com"),
    ("www.acmecorp.com", "acmecorp.com"),
    ("https://acmecorp.com/team/page#frag", "acmecorp.com"),
])
def test_extract_domain_from_url(raw, expected):
    assert extract_domain_from_url(raw) == expected


def test_extract_domain_from_url_empty():
    assert extract_domain_from_url("") == ""


def test_extract_domain_from_url_bare_garbage_best_effort():
    assert extract_domain_from_url("not_a_domain") == "not_a_domain"


def test_extract_domain_from_url_garbage_best_effort_except_branch():
    # urlparse() raises ValueError("Invalid IPv6 URL") -> except fallback path.
    assert extract_domain_from_url("https://[::1") == "[::1"


# ── hunter_email_count ─────────────────────────────────────────────────────────

class _FakeResponse:
    def __init__(self, payload):
        self._payload = payload

    def json(self):
        return self._payload


def test_hunter_email_count_returns_total(monkeypatch):
    captured = {}

    def fake_get(url, params=None, timeout=None):
        captured["url"] = url
        captured["params"] = params
        captured["timeout"] = timeout
        return _FakeResponse({"data": {"total": 7}})

    monkeypatch.setattr(parsers, "get_hunter_api_key", lambda *a, **k: "test-key")
    monkeypatch.setattr("requests.get", fake_get)

    assert hunter_email_count("acmecorp.com") == 7
    assert captured["url"] == "https://api.hunter.io/v2/email-count"
    assert captured["params"] == {"domain": "acmecorp.com", "api_key": "test-key"}
    assert captured["timeout"] == 10


def test_hunter_email_count_empty_domain_makes_no_request(monkeypatch):
    def boom(*args, **kwargs):
        raise AssertionError("requests.get must not be called")

    monkeypatch.setattr("requests.get", boom)
    assert hunter_email_count("") == 0


def test_hunter_email_count_missing_api_key(monkeypatch):
    def boom(*args, **kwargs):
        raise AssertionError("requests.get must not be called")

    monkeypatch.setattr("requests.get", boom)
    monkeypatch.setattr(parsers, "get_hunter_api_key", lambda *a, **k: None)
    assert hunter_email_count("acmecorp.com") == 0


def test_hunter_email_count_request_error_warns(monkeypatch, capsys):
    def boom(*args, **kwargs):
        raise RuntimeError("network down")

    monkeypatch.setattr(parsers, "get_hunter_api_key", lambda *a, **k: "test-key")
    monkeypatch.setattr("requests.get", boom)

    assert hunter_email_count("acmecorp.com") == 0
    assert "[WARN] Hunter email-count error for 'acmecorp.com': network down" in capsys.readouterr().out


# ── parse_shortlister_excel ────────────────────────────────────────────────────

def test_parse_shortlister_excel_extracts_rows(tmp_path):
    path = _write_shortlister_workbook(tmp_path / "report.xlsx", [_holes_row()])
    companies = parse_shortlister_excel(str(path))
    assert companies == [{
        "company_name": "Acme Ltd",
        "ticker": "ACME",
        "industry": "Tech",
        "revenue": "1000",
        "net_income": "100",
        "employees": "50",
        "location": "Mumbai",
        "summary": "Great company",
        "website": "https://acmecorp.com",
        "linkedin_url": "https://linkedin.com/company/acme",
        "linkedin_confirmed": "True",
    }]


def test_parse_shortlister_excel_skips_empty_rows(tmp_path):
    rows = [
        _holes_row(),
        [None] * len(SHORTLISTER_HEADERS),
        _holes_row(**{"Company Name": "Beta Ltd"}),
    ]
    path = _write_shortlister_workbook(tmp_path / "report.xlsx", rows)
    companies = parse_shortlister_excel(str(path))
    assert [c["company_name"] for c in companies] == ["Acme Ltd", "Beta Ltd"]


def test_parse_shortlister_excel_missing_file_warns(tmp_path, capsys):
    assert parse_shortlister_excel(str(tmp_path / "does-not-exist.xlsx")) == []
    assert "[WARN] Could not parse Excel" in capsys.readouterr().out


def test_parse_shortlister_excel_corrupt_file_warns(tmp_path, capsys):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"this is not a real workbook")
    assert parse_shortlister_excel(str(bad)) == []
    assert "[WARN] Could not parse Excel" in capsys.readouterr().out


# ── get_current_session_consolidated_excel ─────────────────────────────────────

def test_get_current_session_consolidated_excel_uses_recorded_last_excel(isolated_var, api_state_reset):
    consolidated = _touch(config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_2020.xlsx")
    shortlister_state.last_excel = str(consolidated)
    assert get_current_session_consolidated_excel() == consolidated


def test_get_current_session_consolidated_excel_prefers_newest_when_last_is_plain(isolated_var, api_state_reset):
    plain = _touch(config.SHORTLISTER_OUTPUT_DIR / "report_2020.xlsx")
    _touch(config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_old.xlsx", mtime=1_000_000)
    newer = _touch(config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_new.xlsx", mtime=2_000_000)
    shortlister_state.last_excel = str(plain)
    assert get_current_session_consolidated_excel() == newer


def test_get_current_session_consolidated_excel_none_when_nothing_exists(isolated_var, api_state_reset):
    assert get_current_session_consolidated_excel() is None


def test_get_current_session_consolidated_excel_ignores_temp_files(isolated_var, api_state_reset):
    _touch(config.SHORTLISTER_OUTPUT_DIR / "~$report_consolidated_2020.xlsx")
    assert get_current_session_consolidated_excel() is None


# ── get_all_decision_makers ───────────────────────────────────────────────────

def test_get_all_decision_makers_flat_layout(isolated_var, api_state_reset):
    path = config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_flat.xlsx"
    _write_consolidated_flat(path, [
        ["Acme Ltd", "https://acmecorp.com", "https://linkedin.com/company/acme", "Mumbai",
         "CEO", "Jane Doe", "https://linkedin.com/in/jane"],
        ["Beta Ltd", "https://beta.com", "", "Delhi",
         "CTO", "John Roe", "https://linkedin.com/in/john"],
    ])
    shortlister_state.last_excel = str(path)

    dms = get_all_decision_makers()
    assert dms == [
        {
            "company_name": "Acme Ltd",
            "company_website": "https://acmecorp.com",
            "company_linkedin": "https://linkedin.com/company/acme",
            "location": "Mumbai",
            "name": "Jane Doe",
            "position": "CEO",
            "linkedin_url": "https://linkedin.com/in/jane",
            "source_file": "report_consolidated_flat.xlsx",
            "id": "dm-1",
        },
        {
            "company_name": "Beta Ltd",
            "company_website": "https://beta.com",
            "company_linkedin": "",
            "location": "Delhi",
            "name": "John Roe",
            "position": "CTO",
            "linkedin_url": "https://linkedin.com/in/john",
            "source_file": "report_consolidated_flat.xlsx",
            "id": "dm-2",
        },
    ]


def test_get_all_decision_makers_collapsible_layout(isolated_var, api_state_reset):
    path = config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_grouped.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Decision Maker Contacts"
    ws.append(["  ▸ Acme Ltd  (2 contacts)"])
    ws.append(["#", "Name", "Position", "LinkedIn (Person)", "Company LinkedIn", "Location"])
    ws.append([1, "Jane Doe", "CEO", "https://linkedin.com/in/jane",
               "https://linkedin.com/company/acme", "Mumbai"])
    ws.append([2, "—", "CFO", "https://linkedin.com/in/skipped", None, "Delhi"])
    wb.save(path)
    shortlister_state.last_excel = str(path)

    dms = get_all_decision_makers()
    assert dms == [{
        "company_name": "Acme Ltd",
        "company_website": "",
        "company_linkedin": "https://linkedin.com/company/acme",
        "location": "Mumbai",
        "name": "Jane Doe",
        "position": "CEO",
        "linkedin_url": "https://linkedin.com/in/jane",
        "source_file": "report_consolidated_grouped.xlsx",
        "id": "dm-1",
    }]


def test_get_all_decision_makers_dedup_keeps_first(isolated_var, api_state_reset):
    path = config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_dup.xlsx"
    _write_consolidated_flat(path, [
        ["Acme Ltd", "https://acmecorp.com", "", "Mumbai", "CEO", "Jane Doe", "https://linkedin.com/in/first"],
        ["Acme Ltd", "https://acmecorp.com", "", "Mumbai", "COO", "jane doe", "https://linkedin.com/in/second"],
    ])
    shortlister_state.last_excel = str(path)

    dms = get_all_decision_makers()
    assert len(dms) == 1
    assert dms[0]["position"] == "CEO"
    assert dms[0]["linkedin_url"] == "https://linkedin.com/in/first"
    assert dms[0]["id"] == "dm-1"


def test_get_all_decision_makers_json_fallback(isolated_var, api_state_reset):
    reports = config.SHORTLISTER_OUTPUT_DIR
    reports.mkdir(parents=True, exist_ok=True)
    (reports / "report_session.json").write_text(json.dumps({
        "companies": [{
            "name": "Acme Ltd",
            "website": "https://acmecorp.com",
            "linkedin_verified": "https://linkedin.com/company/acme",
            "location": "Mumbai",
            "decision_makers": [
                {"name": "Jane Doe", "position": "CEO", "linkedin_url": "https://linkedin.com/in/jane"},
                {"name": "—", "position": "CFO", "linkedin_url": "https://linkedin.com/in/skip"},
            ],
        }]
    }), encoding="utf-8")

    dms = get_all_decision_makers()
    assert dms == [{
        "company_name": "Acme Ltd",
        "company_website": "https://acmecorp.com",
        "company_linkedin": "https://linkedin.com/company/acme",
        "location": "Mumbai",
        "name": "Jane Doe",
        "position": "CEO",
        "linkedin_url": "https://linkedin.com/in/jane",
        "source_file": "report_session.json",
        "id": "dm-1",
    }]


def test_get_all_decision_makers_cross_references_linkedin_outputs(isolated_var, api_state_reset):
    path = config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_cross.xlsx"
    _write_consolidated_flat(path, [
        ["Acme Ltd", "https://acmecorp.com", "", "Mumbai", "CEO", "Jane Doe", ""],
    ])
    shortlister_state.last_excel = str(path)

    linkedin_dir = config.LINKEDIN_OUTPUT_DIR
    linkedin_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(["Company Name", "X", "Y", "Z", "W", "Decision Maker Name", "Decision Maker LinkedIn"])
    ws.append(["Acme Ltd", "", "", "", "", "Jane Doe", "https://linkedin.com/in/from-contacts"])
    wb.save(linkedin_dir / "contacts_acme.xlsx")

    dms = get_all_decision_makers()
    assert len(dms) == 1
    assert dms[0]["linkedin_url"] == "https://linkedin.com/in/from-contacts"
    assert dms[0]["id"] == "dm-1"


# ── extract_location_from_query ────────────────────────────────────────────────

@pytest.mark.asyncio
@pytest.mark.parametrize("query, expected", [
    ("indian companies", "India"),
    ("companies in india", "India"),
    ("top us firms", "United States"),
    ("united states market", "United States"),
    ("american companies", "United States"),
    ("companies in the uk", "United Kingdom"),
    ("united kingdom firms", "United Kingdom"),
    ("british companies", "United Kingdom"),
    ("canada exporters", "Canada"),
    ("canadian firms", "Canada"),
    ("germany manufacturers", "Germany"),
    ("german companies", "Germany"),
])
async def test_extract_location_from_query_keywords(query, expected):
    assert await extract_location_from_query(query) == expected


@pytest.mark.asyncio
async def test_extract_location_from_query_empty_defaults_to_india():
    assert await extract_location_from_query("") == "India"
    assert await extract_location_from_query("   ") == "India"


def _patch_openai_config(monkeypatch):
    monkeypatch.setattr(config, "OPENAI_BASE_URL", "https://llm.example.test/v1")
    monkeypatch.setattr(config, "OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(config, "OPENAI_MODEL", "test-model")


class _FakeLLMResult:
    def __init__(self, content):
        self.content = content


@pytest.mark.asyncio
async def test_extract_location_from_query_llm_fallback_strips_quotes(monkeypatch):
    _patch_openai_config(monkeypatch)
    captured = {}

    class FakeChat:
        def __init__(self, **kwargs):
            captured.update(kwargs)

        async def ainvoke(self, prompt):
            captured["prompt"] = prompt
            return _FakeLLMResult('  "France"  ')

    monkeypatch.setattr("langchain_openai.ChatOpenAI", FakeChat)

    assert await extract_location_from_query("firms in france") == "France"
    assert captured["model"] == "test-model"
    assert captured["base_url"] == "https://llm.example.test/v1"
    assert "firms in france" in captured["prompt"]


@pytest.mark.asyncio
async def test_extract_location_from_query_llm_error_defaults_to_india(monkeypatch, capsys):
    _patch_openai_config(monkeypatch)

    class BoomChat:
        def __init__(self, **kwargs):
            pass

        async def ainvoke(self, prompt):
            raise RuntimeError("llm down")

    monkeypatch.setattr("langchain_openai.ChatOpenAI", BoomChat)

    assert await extract_location_from_query("firms in france") == "India"
    assert "[WARN] Location extraction error: llm down" in capsys.readouterr().out


@pytest.mark.asyncio
async def test_extract_location_from_query_without_llm_config_defaults_to_india(monkeypatch):
    monkeypatch.setattr(config, "OPENAI_BASE_URL", None)
    monkeypatch.setattr(config, "OPENAI_API_KEY", None)
    monkeypatch.setattr(config, "OPENAI_MODEL", None)

    assert await extract_location_from_query("firms in france") == "India"
