"""Unit tests for ``leadgen.api.runner`` — broadcast/sentinel fan-out, subprocess
streaming, report discovery, the background runners and the SSE generator."""

import asyncio
import json
import os
import sys
from unittest.mock import AsyncMock

import pytest
from openpyxl import Workbook

from leadgen import config
from leadgen.api import runner
from leadgen.api.runner import (
    broadcast,
    make_sse_generator,
    newest_shortlister_report,
    run_linkedin_bg,
    run_shortlister_bg,
    send_sentinel,
    stream_subprocess,
)
from leadgen.api.state import AgentState, linkedin_state, shortlister_state, ws_log_clients

SHORTLISTER_HEADERS = [
    "Company Name", "Ticker", "Industry", "Revenue (TTM)", "Net Income (TTM)",
    "Employee Count", "Location", "Business Summary", "Website",
    "LinkedIn (Verified)", "LinkedIn Confirmed?",
]

SHORTLISTER_ROW = [
    "Acme Ltd", "ACME", "Tech", 1000, 100, 50, "Mumbai", "Summary",
    "https://acmecorp.com", "linkedin.com/company/acme", True,
]


class FakeWebSocket:
    def __init__(self, fail=False):
        self.sent = []
        self.fail = fail

    async def send_text(self, text):
        if self.fail:
            raise RuntimeError("socket closed")
        self.sent.append(text)


class FakeRequest:
    async def is_disconnected(self):
        return False


def _touch(path, mtime=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"x")
    if mtime is not None:
        os.utime(path, (mtime, mtime))
    return path


def _write_report(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    for ci, header in enumerate(SHORTLISTER_HEADERS, 1):
        ws.cell(row=1, column=ci, value=header)
    for ci, value in enumerate(SHORTLISTER_ROW, 1):
        ws.cell(row=2, column=ci, value=value)
    wb.save(path)
    return path


def _drain(queue_):
    items = []
    while not queue_.empty():
        items.append(queue_.get_nowait())
    return items


# ── broadcast ──────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_broadcast_buffers_and_delivers_to_subscribers(api_state_reset):
    state = AgentState()
    q = asyncio.Queue()
    state.log_queues.append(q)

    await broadcast(state, "hello world")

    assert state.log_buffer == ["hello world"]
    assert q.get_nowait() == "hello world"


@pytest.mark.asyncio
async def test_broadcast_removes_full_queue(api_state_reset):
    state = AgentState()
    full = asyncio.Queue(maxsize=1)
    full.put_nowait("occupied")
    state.log_queues.append(full)

    await broadcast(state, "line")

    assert full not in state.log_queues
    assert state.log_buffer == ["line"]


@pytest.mark.asyncio
async def test_broadcast_fans_out_to_ws_clients_and_prunes_failures(api_state_reset):
    state = AgentState()
    good = FakeWebSocket()
    bad = FakeWebSocket(fail=True)
    ws_log_clients.extend([good, bad])

    await broadcast(state, "payload")

    envelope = json.dumps({"event_type": "log", "data": {"message": "payload"}})
    assert good.sent == [envelope]
    assert good in ws_log_clients
    assert bad not in ws_log_clients


# ── send_sentinel ──────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_send_sentinel_puts_none_into_every_queue(api_state_reset):
    state = AgentState()
    q1, q2 = asyncio.Queue(), asyncio.Queue()
    state.log_queues.extend([q1, q2])

    await send_sentinel(state)

    assert q1.get_nowait() is None
    assert q2.get_nowait() is None


@pytest.mark.asyncio
async def test_send_sentinel_ignores_full_queue(api_state_reset):
    state = AgentState()
    full = asyncio.Queue(maxsize=1)
    full.put_nowait("x")
    state.log_queues.append(full)

    await send_sentinel(state)

    assert full.get_nowait() == "x"


# ── stream_subprocess ──────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_stream_subprocess_collects_lines(api_state_reset):
    state = AgentState()
    rc = await stream_subprocess(
        state,
        [sys.executable, "-u", "-c", "print('alpha'); print('beta')"],
        config.ROOT_DIR,
    )
    assert rc == 0
    assert state.log_buffer == ["alpha", "beta"]


@pytest.mark.asyncio
async def test_stream_subprocess_applies_prefix(api_state_reset):
    state = AgentState()
    rc = await stream_subprocess(
        state,
        [sys.executable, "-u", "-c", "print('alpha'); print('beta')"],
        config.ROOT_DIR,
        prefix="X",
    )
    assert rc == 0
    assert state.log_buffer == ["[X] alpha", "[X] beta"]


@pytest.mark.asyncio
async def test_stream_subprocess_writes_stdin(api_state_reset):
    state = AgentState()
    rc = await stream_subprocess(
        state,
        [sys.executable, "-u", "-c", "import sys; print(sys.stdin.read().strip())"],
        config.ROOT_DIR,
        stdin_input="hello from stdin\n",
    )
    assert rc == 0
    assert state.log_buffer == ["hello from stdin"]


# ── newest_shortlister_report ──────────────────────────────────────────────────

def test_newest_shortlister_report_prefers_consolidated(isolated_var):
    out = config.SHORTLISTER_OUTPUT_DIR
    _touch(out / "report_2024.xlsx", mtime=3_000_000)
    consolidated = _touch(out / "report_consolidated_2020.xlsx", mtime=1_000_000)
    assert newest_shortlister_report() == consolidated


def test_newest_shortlister_report_none_for_missing_dir(isolated_var):
    assert newest_shortlister_report() is None


def test_newest_shortlister_report_ignores_temp_files(isolated_var):
    _touch(config.SHORTLISTER_OUTPUT_DIR / "~$report_consolidated_2020.xlsx")
    assert newest_shortlister_report() is None


# ── run_shortlister_bg ─────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_run_shortlister_bg_success(isolated_var, api_state_reset, monkeypatch):
    report = _write_report(config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_run.xlsx")

    async def fake_stream(state, cmd, cwd, stdin_input=None, prefix=None, env_overrides=None):
        return 0

    async def fake_location(query):
        return "India"

    monkeypatch.setattr(runner, "stream_subprocess", fake_stream)
    monkeypatch.setattr(runner, "extract_location_from_query", fake_location)
    done = AsyncMock()
    monkeypatch.setattr(runner.visualizer_broker, "mark_process_done", done)

    q = asyncio.Queue()
    shortlister_state.log_queues.append(q)

    await run_shortlister_bg("top software companies", "run-123")

    assert shortlister_state.status == "done"
    assert shortlister_state.user_query == "top software companies"
    assert shortlister_state.extracted_location == "India"
    assert shortlister_state.last_excel == str(report)
    assert [c["company_name"] for c in shortlister_state.companies] == ["Acme Ltd"]
    assert shortlister_state.error is None

    done.assert_awaited_once_with("run-123", "done", None)
    assert None in _drain(q)


@pytest.mark.asyncio
async def test_run_shortlister_bg_failure_sets_error(isolated_var, api_state_reset, monkeypatch):
    async def fake_stream(state, cmd, cwd, stdin_input=None, prefix=None, env_overrides=None):
        return 1

    async def fake_location(query):
        return "India"

    monkeypatch.setattr(runner, "stream_subprocess", fake_stream)
    monkeypatch.setattr(runner, "extract_location_from_query", fake_location)
    done = AsyncMock()
    monkeypatch.setattr(runner.visualizer_broker, "mark_process_done", done)

    await run_shortlister_bg("query", "run-1")

    assert shortlister_state.status == "error"
    assert shortlister_state.error == "Process exited with code 1"
    done.assert_awaited_once_with("run-1", "error", "Process exited with code 1")


# ── run_linkedin_bg ────────────────────────────────────────────────────────────

def _company(name, linkedin="", website="", location="India"):
    return {"company_name": name, "linkedin_url": linkedin, "website": website, "location": location}


@pytest.mark.asyncio
async def test_run_linkedin_bg_builds_argv(api_state_reset, monkeypatch):
    calls = []

    async def fake_stream(state, cmd, cwd, stdin_input=None, prefix=None, env_overrides=None):
        calls.append(cmd)
        return 0

    monkeypatch.setattr(runner, "stream_subprocess", fake_stream)

    companies = [
        _company("Acme", linkedin="https://linkedin.com/company/acme", website="https://acme.com"),
        _company("Beta", linkedin="NOT FOUND", website="https://beta.com"),
    ]
    await run_linkedin_bg(companies, concurrency=2)

    assert linkedin_state.status == "done"
    assert linkedin_state.total_count == 2
    assert linkedin_state.completed_count == 2
    assert linkedin_state.error is None

    acme = next(cmd for cmd in calls if "Acme" in cmd)
    assert "-m" in acme
    assert "leadgen.agents.linkedin_finder" in acme
    assert acme[acme.index("--company_name") + 1] == "Acme"
    assert acme[acme.index("--company_website") + 1] == "https://acme.com"
    assert acme[acme.index("--location") + 1] == "India"
    assert acme[acme.index("--company_linkedin") + 1] == "https://linkedin.com/company/acme"

    beta = next(cmd for cmd in calls if "Beta" in cmd)
    assert "--company_linkedin" not in beta
    assert beta[beta.index("--company_website") + 1] == "https://beta.com"


@pytest.mark.asyncio
async def test_run_linkedin_bg_uses_shortlister_location_fallback(api_state_reset, monkeypatch):
    calls = []

    async def fake_stream(state, cmd, cwd, stdin_input=None, prefix=None, env_overrides=None):
        calls.append(cmd)
        return 0

    monkeypatch.setattr(runner, "stream_subprocess", fake_stream)
    shortlister_state.extracted_location = "Canada"

    await run_linkedin_bg([{"company_name": "Acme", "website": ""}], concurrency=1)

    acme = next(cmd for cmd in calls if "Acme" in cmd)
    assert acme[acme.index("--location") + 1] == "Canada"


@pytest.mark.asyncio
async def test_run_linkedin_bg_all_failures_set_error(api_state_reset, monkeypatch):
    async def fake_stream(state, cmd, cwd, stdin_input=None, prefix=None, env_overrides=None):
        return 1

    monkeypatch.setattr(runner, "stream_subprocess", fake_stream)
    companies = [_company("Acme"), _company("Beta")]
    await run_linkedin_bg(companies, concurrency=1)

    assert linkedin_state.status == "error"
    assert linkedin_state.error == (
        "Completed with errors in 2 company(ies): 'Acme' (exit code 1), 'Beta' (exit code 1)"
    )
    assert linkedin_state.completed_count == 2


@pytest.mark.asyncio
async def test_run_linkedin_bg_partial_failure_is_done_with_error(api_state_reset, monkeypatch):
    async def fake_stream(state, cmd, cwd, stdin_input=None, prefix=None, env_overrides=None):
        return 0 if cmd[cmd.index("--company_name") + 1] == "Acme" else 1

    monkeypatch.setattr(runner, "stream_subprocess", fake_stream)
    companies = [_company("Acme"), _company("Beta")]
    await run_linkedin_bg(companies, concurrency=1)

    assert linkedin_state.status == "done"
    assert linkedin_state.error == "Completed with errors in 1 company(ies): 'Beta' (exit code 1)"
    assert linkedin_state.completed_count == 2


@pytest.mark.asyncio
async def test_run_linkedin_bg_respects_concurrency_semaphore(api_state_reset, monkeypatch):
    active = 0
    max_active = 0

    async def fake_stream(state, cmd, cwd, stdin_input=None, prefix=None, env_overrides=None):
        nonlocal active, max_active
        active += 1
        max_active = max(max_active, active)
        await asyncio.sleep(0.02)
        active -= 1
        return 0

    monkeypatch.setattr(runner, "stream_subprocess", fake_stream)
    companies = [_company(f"C{i}") for i in range(4)]
    await run_linkedin_bg(companies, concurrency=1)

    assert max_active == 1
    assert linkedin_state.completed_count == 4
    assert linkedin_state.status == "done"
    assert linkedin_state.active_companies == []


# ── make_sse_generator ─────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_make_sse_generator_idle_yields_nothing(api_state_reset):
    state = AgentState()
    state.status = "idle"
    gen = make_sse_generator(state, FakeRequest())()
    assert [item async for item in gen] == []


@pytest.mark.asyncio
async def test_make_sse_generator_done_replays_buffer(api_state_reset):
    state = AgentState()
    state.status = "done"
    state.log_buffer.extend(["a", "b"])
    gen = make_sse_generator(state, FakeRequest())()
    assert [item async for item in gen] == [
        {"event": "log", "data": "a"},
        {"event": "log", "data": "b"},
        {"event": "done", "data": "done"},
    ]


@pytest.mark.asyncio
async def test_make_sse_generator_running_streams_live(api_state_reset):
    state = AgentState()
    state.status = "running"
    gen = make_sse_generator(state, FakeRequest())()

    async def pusher():
        await asyncio.sleep(0.01)
        await broadcast(state, "live-line")
        await send_sentinel(state)

    task = asyncio.create_task(pusher())
    items = [item async for item in gen]
    await task

    assert items == [
        {"event": "log", "data": "live-line"},
        {"event": "done", "data": "running"},
    ]
    with pytest.raises(StopAsyncIteration):
        await gen.__anext__()
