"""End-to-end integration tests over the FastAPI router surface via TestClient."""

import re

import pytest
from openpyxl import Workbook
from starlette.testclient import TestClient

from leadgen import config
from leadgen.api.app import app
from leadgen.api.state import _email_store, linkedin_state, shortlister_state

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FLAT_DM_HEADERS = [
    "Company Name", "Company Website", "Company LinkedIn", "Location",
    "Decision Maker Position", "Decision Maker Name", "Decision Maker LinkedIn",
]


def _write_consolidated_flat(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Decision Maker Contacts"
    for ci, header in enumerate(FLAT_DM_HEADERS, 1):
        ws.cell(row=1, column=ci, value=header)
    for ri, row in enumerate(rows, 2):
        for ci, value in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=value)
    wb.save(path)
    return path


def _dm(dm_id, name, company, website, linkedin):
    return {
        "id": dm_id, "name": name, "company_name": company,
        "company_website": website, "linkedin_url": linkedin,
    }


# ── health / status shapes ─────────────────────────────────────────────────────

def test_health(isolated_var, api_state_reset):
    with TestClient(app) as client:
        resp = client.get("/health")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


def test_shortlister_status_shape(isolated_var, api_state_reset):
    with TestClient(app) as client:
        resp = client.get("/shortlister/status")
    assert resp.status_code == 200
    assert resp.json() == {
        "status": "idle",
        "excel_path": None,
        "companies": [],
        "error": None,
    }


def test_linkedin_status_shape(isolated_var, api_state_reset):
    with TestClient(app) as client:
        resp = client.get("/linkedin/status")
    assert resp.status_code == 200
    assert resp.json() == {
        "status": "idle",
        "current_company": None,
        "active_companies": [],
        "completed_count": 0,
        "total_count": 0,
        "error": None,
    }


# ── linkedin files / download ──────────────────────────────────────────────────

def test_linkedin_files_lists_and_skips_temp(isolated_var, api_state_reset):
    out = config.LINKEDIN_OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    (out / "contacts_acme.xlsx").write_bytes(b"data")
    (out / "~$contacts_acme.xlsx").write_bytes(b"lock")

    with TestClient(app) as client:
        resp = client.get("/linkedin/files")

    assert resp.status_code == 200
    files = resp.json()["files"]
    assert [f["filename"] for f in files] == ["contacts_acme.xlsx"]
    assert files[0]["download_url"] == "/linkedin/download/contacts_acme.xlsx"
    assert files[0]["size_bytes"] == len(b"data")


def test_linkedin_download_ok_and_missing(isolated_var, api_state_reset):
    out = config.LINKEDIN_OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    (out / "contacts_acme.xlsx").write_bytes(b"xlsx-bytes")

    with TestClient(app) as client:
        ok = client.get("/linkedin/download/contacts_acme.xlsx")
        missing = client.get("/linkedin/download/missing.xlsx")

    assert ok.status_code == 200
    assert ok.headers["content-type"] == XLSX_MEDIA_TYPE
    assert ok.content == b"xlsx-bytes"
    assert missing.status_code == 404


def test_linkedin_download_path_traversal_neutralised(isolated_var, api_state_reset):
    runs = config.RUNS_DIR
    runs.mkdir(parents=True, exist_ok=True)
    (runs / "secrets.xlsx").write_bytes(b"top-secret")

    with TestClient(app) as client:
        encoded = client.get("/linkedin/download/..%2F..%2Fsecrets.xlsx")
        plain = client.get("/linkedin/download/../secrets.xlsx")

    assert encoded.status_code == 404
    assert plain.status_code == 404


@pytest.mark.asyncio
async def test_linkedin_download_traversal_neutralised_by_basename(isolated_var, api_state_reset):
    from fastapi import HTTPException

    from leadgen.api.routers import linkedin as linkedin_router

    runs = config.RUNS_DIR
    runs.mkdir(parents=True, exist_ok=True)
    (runs / "secrets.xlsx").write_bytes(b"top-secret")

    with pytest.raises(HTTPException) as excinfo:
        await linkedin_router.linkedin_download_file("../secrets.xlsx")
    assert excinfo.value.status_code == 404


# ── exports ────────────────────────────────────────────────────────────────────

def test_export_list_empty_then_populated(isolated_var, api_state_reset):
    with TestClient(app) as client:
        empty = client.get("/export/list")
        config.EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        (config.EXPORTS_DIR / "export_2020-01-01_00-00-00.xlsx").write_bytes(b"x")
        populated = client.get("/export/list")

    assert empty.status_code == 200
    assert empty.json() == {"files": []}
    files = populated.json()["files"]
    assert [f["filename"] for f in files] == ["export_2020-01-01_00-00-00.xlsx"]
    assert files[0]["download_url"] == "/export/download/export_2020-01-01_00-00-00.xlsx"


def test_export_download_ok_and_missing(isolated_var, api_state_reset):
    config.EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    (config.EXPORTS_DIR / "export_2020-01-01_00-00-00.xlsx").write_bytes(b"export-bytes")

    with TestClient(app) as client:
        ok = client.get("/export/download/export_2020-01-01_00-00-00.xlsx")
        missing = client.get("/export/download/nope.xlsx")

    assert ok.status_code == 200
    assert ok.content == b"export-bytes"
    assert missing.status_code == 404


def test_export_consolidated_streams_bytes_and_writes_copy(isolated_var, api_state_reset, monkeypatch):
    from leadgen.api.routers import exports as exports_router

    monkeypatch.setattr(exports_router, "build_export_xlsx", lambda *args, **kwargs: b"fake-xlsx")

    with TestClient(app) as client:
        resp = client.get("/export/consolidated")

    assert resp.status_code == 200
    assert resp.content == b"fake-xlsx"
    disposition = resp.headers["content-disposition"]
    assert re.fullmatch(r'attachment; filename="export_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.xlsx"', disposition)

    written = list(config.EXPORTS_DIR.glob("export_*.xlsx"))
    assert len(written) == 1
    assert written[0].read_bytes() == b"fake-xlsx"


# ── reset ──────────────────────────────────────────────────────────────────────

def test_reset_clears_agent_state(isolated_var, api_state_reset):
    shortlister_state.status = "running"
    shortlister_state.companies = [{"company_name": "Acme Ltd"}]
    linkedin_state.status = "running"
    linkedin_state.completed_count = 5

    with TestClient(app) as client:
        resp = client.post("/reset")

    assert resp.status_code == 200
    assert resp.json() == {"status": "reset"}
    assert shortlister_state.status == "idle"
    assert shortlister_state.companies == []
    assert linkedin_state.status == "idle"
    assert linkedin_state.completed_count == 0


# ── shortlister download ───────────────────────────────────────────────────────

def test_shortlister_download_404_then_200(isolated_var, api_state_reset):
    with TestClient(app) as client:
        missing = client.get("/shortlister/download")
        xlsx = config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_1.xlsx"
        xlsx.parent.mkdir(parents=True, exist_ok=True)
        xlsx.write_bytes(b"report-bytes")
        shortlister_state.last_excel = str(xlsx)
        ok = client.get("/shortlister/download")

    assert missing.status_code == 404
    assert ok.status_code == 200
    assert ok.headers["content-type"] == XLSX_MEDIA_TYPE
    assert ok.content == b"report-bytes"


# ── decision makers ────────────────────────────────────────────────────────────

def test_linkedin_decision_makers_parses_consolidated(isolated_var, api_state_reset):
    path = _write_consolidated_flat(config.SHORTLISTER_OUTPUT_DIR / "report_consolidated_1.xlsx", [
        ["Acme Ltd", "https://acmecorp.com", "https://linkedin.com/company/acme", "Mumbai",
         "CEO", "Jane Doe", "https://linkedin.com/in/jane"],
    ])
    shortlister_state.last_excel = str(path)

    with TestClient(app) as client:
        resp = client.get("/linkedin/decision-makers")

    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    assert body["decision_makers"] == [{
        "company_name": "Acme Ltd",
        "company_website": "https://acmecorp.com",
        "company_linkedin": "https://linkedin.com/company/acme",
        "location": "Mumbai",
        "name": "Jane Doe",
        "position": "CEO",
        "linkedin_url": "https://linkedin.com/in/jane",
        "source_file": "report_consolidated_1.xlsx",
        "id": "dm-1",
    }]


def test_linkedin_decision_makers_empty_when_nothing_exists(isolated_var, api_state_reset):
    with TestClient(app) as client:
        resp = client.get("/linkedin/decision-makers")

    assert resp.status_code == 200
    assert resp.json() == {"decision_makers": [], "total": 0}


# ── hunter: store-email / email-count / find-email ─────────────────────────────

def test_hunter_store_email_lowercases_key(isolated_var, api_state_reset):
    with TestClient(app) as client:
        resp = client.post("/hunter/store-email", json={
            "company_name": "Acme Ltd", "dm_name": "Jane Doe", "email": "Jane@Acme.com",
        })

    assert resp.status_code == 200
    assert resp.json() == {"stored": True, "key": "Acme Ltd / Jane Doe"}
    assert _email_store[("acme ltd", "jane doe")] == "Jane@Acme.com"


def test_hunter_store_email_blank_rejected(isolated_var, api_state_reset):
    with TestClient(app) as client:
        resp = client.post("/hunter/store-email", json={
            "company_name": "Acme Ltd", "dm_name": "Jane Doe", "email": "   ",
        })
    assert resp.status_code == 400
    assert _email_store == {}


def test_hunter_email_count_endpoint(isolated_var, api_state_reset, monkeypatch):
    from leadgen.api.routers import hunter as hunter_router

    monkeypatch.setattr(hunter_router, "hunter_email_count", lambda domain: 5)

    with TestClient(app) as client:
        resp = client.get("/hunter/email-count", params={"domain": "acme.com"})

    assert resp.status_code == 200
    assert resp.json() == {"domain": "acme.com", "count": 5}


def test_hunter_find_email_returns_payload(isolated_var, api_state_reset, monkeypatch):
    from leadgen.api.routers import hunter as hunter_router

    payload = {"data": {"email": "jane@acme.com", "score": 95}}
    monkeypatch.setattr(hunter_router, "find_email_by_linkedin", lambda **kwargs: payload)

    with TestClient(app) as client:
        resp = client.post("/hunter/find-email", json={"linkedin_url": "https://linkedin.com/in/jane"})

    assert resp.status_code == 200
    assert resp.json() == payload


def test_hunter_find_email_blank_linkedin_rejected(isolated_var, api_state_reset):
    with TestClient(app) as client:
        resp = client.post("/hunter/find-email", json={"linkedin_url": "   "})
    assert resp.status_code == 400


def test_hunter_find_email_value_error_rejected(isolated_var, api_state_reset, monkeypatch):
    from leadgen.api.routers import hunter as hunter_router

    def boom(**kwargs):
        raise ValueError("bad handle")

    monkeypatch.setattr(hunter_router, "find_email_by_linkedin", boom)

    with TestClient(app) as client:
        resp = client.post("/hunter/find-email", json={"linkedin_url": "handle"})

    assert resp.status_code == 400
    assert resp.json()["detail"] == "bad handle"


def test_hunter_find_email_generic_error_is_500(isolated_var, api_state_reset, monkeypatch):
    from leadgen.api.routers import hunter as hunter_router

    def boom(**kwargs):
        raise RuntimeError("kaboom")

    monkeypatch.setattr(hunter_router, "find_email_by_linkedin", boom)

    with TestClient(app) as client:
        resp = client.post("/hunter/find-email", json={"linkedin_url": "handle"})

    assert resp.status_code == 500
    assert resp.json()["detail"] == "Hunter API error: kaboom"


def test_hunter_smart_enrich_caps_skips_and_filters(isolated_var, api_state_reset, monkeypatch):
    from leadgen.api.routers import hunter as hunter_router

    counts = {"acme.com": 3, "beta.com": 0, "gamma.com": 5}
    monkeypatch.setattr(hunter_router, "hunter_email_count", lambda domain: counts.get(domain, 0))

    def fake_finder(**kwargs):
        return {"data": {"email": "found@acme.com", "score": 80}}

    monkeypatch.setattr(hunter_router, "find_email_by_linkedin", fake_finder)

    decision_makers = [
        _dm(f"a{i}", f"Person {i}", "Acme", "https://acme.com", "https://linkedin.com/in/a") for i in range(1, 6)
    ]
    decision_makers += [
        _dm("b1", "Beta One", "Beta", "https://beta.com", "https://linkedin.com/in/b1"),
        _dm("b2", "Beta Two", "Beta", "https://beta.com", "https://linkedin.com/in/b2"),
    ]
    decision_makers += [
        _dm("c1", "Gamma One", "Gamma", "https://gamma.com", None),
    ]

    with TestClient(app) as client:
        resp = client.post("/hunter/smart-enrich", json={"decision_makers": decision_makers})

    assert resp.status_code == 200
    body = resp.json()
    assert set(body.keys()) == {
        "enriched", "no_indexed_dm_ids", "skipped_companies",
        "stopped_at_cap_companies", "total_emails_found",
    }
    assert [e["id"] for e in body["enriched"]] == ["a1", "a2", "a3"]
    assert all(e["email"] == "found@acme.com" for e in body["enriched"])
    assert sorted(body["no_indexed_dm_ids"]) == ["b1", "b2"]
    assert body["skipped_companies"] == ["Beta"]
    assert body["stopped_at_cap_companies"] == ["Acme"]
    assert body["total_emails_found"] == 3
