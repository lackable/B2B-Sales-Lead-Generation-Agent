"""
Agent Pipeline Backend
======================
Wraps Shortlister Agent and LinkedIn Finder Agent via subprocess execution.
Stdout is teed to the terminal (unchanged) and broadcast over SSE to the browser.
"""

import asyncio
import io
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import json
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
from fastapi import BackgroundTasks, FastAPI, HTTPException, Request, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from sse_starlette.sse import EventSourceResponse

# ── Directory Paths ────────────────────────────────────────────────────────────

BASE_DIR: Path = Path(__file__).parent.parent.resolve()
SHORTLISTER_DIR: Path = BASE_DIR / "Shortlister Agent"
LINKEDIN_DIR: Path = BASE_DIR / "Linkedin Finder Agent"
PARENT_OUTPUT_DIR: Path = BASE_DIR / "output"

# Add BASE_DIR to sys.path to import hunter_linkedin_email_finder
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

try:
    from hunter_linkedin_email_finder import find_email_by_linkedin, get_hunter_api_key
except ImportError:
    find_email_by_linkedin = None
    get_hunter_api_key = None

# Create parent output directories at startup
(PARENT_OUTPUT_DIR / "shortlister").mkdir(parents=True, exist_ok=True)
(PARENT_OUTPUT_DIR / "linkedin").mkdir(parents=True, exist_ok=True)
(PARENT_OUTPUT_DIR / "exports").mkdir(parents=True, exist_ok=True)

# ── Shared WebSocket log broadcaster ──────────────────────────────────────────
# The logging_utils.ws_broadcaster.ws_queue singleton is populated by
# WebSocketQueueHandler whenever an agent calls logger.event(). The
# /ws/logs endpoint drains this queue and pushes JSON to the browser.
try:
    from logging_utils.ws_broadcaster import ws_queue as _ws_log_queue
except ImportError:
    import asyncio as _asyncio
    _ws_log_queue = _asyncio.Queue(maxsize=2000)

# Connected WebSocket clients for structured log streaming
_ws_log_clients: list = []

try:
    from export_builder import build_export_xlsx
except ImportError:
    build_export_xlsx = None  # type: ignore

try:
    from visualizer_broker import VisualizerBroker
except ImportError:
    from backend.visualizer_broker import VisualizerBroker

# ── In-Memory Email Store ──────────────────────────────────────────────────────
# Keyed by (company_name_lower, dm_name_lower) → email string.
# Populated in real-time whenever Hunter finds an email in the frontend.
_email_store: Dict[tuple, str] = {}


# ── Shared Agent State ─────────────────────────────────────────────────────────

class AgentState:
    def __init__(self):
        self.status: str = "idle"           # idle | running | done | error
        self.log_queues: List[asyncio.Queue] = []
        self.log_buffer: List[str] = []     # replay buffer for reconnecting clients
        self.last_excel: Optional[str] = None
        self.companies: List[Dict] = []
        self.current_company: Optional[str] = None
        self.active_companies: List[str] = []
        self.completed_count: int = 0
        self.total_count: int = 0
        self.user_query: str = ""
        self.extracted_location: str = ""
        self.error: Optional[str] = None

    def reset(self):
        self.status = "idle"
        self.log_queues.clear()
        self.log_buffer.clear()
        self.last_excel = None
        self.companies.clear()
        self.current_company = None
        self.active_companies.clear()
        self.completed_count = 0
        self.total_count = 0
        self.user_query = ""
        self.extracted_location = ""
        self.error = None


shortlister_state = AgentState()
linkedin_state = AgentState()


async def _extract_location_from_query(query: str) -> str:
    """Extract location string from user's Shortlister query."""
    if not query or not query.strip():
        return "India"
    
    q_lower = query.lower()
    if "indian" in q_lower or "india" in q_lower:
        return "India"
    if "us" in q_lower or "united states" in q_lower or "american" in q_lower:
        return "United States"
    if "uk" in q_lower or "united kingdom" in q_lower or "british" in q_lower:
        return "United Kingdom"
    if "canada" in q_lower or "canadian" in q_lower:
        return "Canada"
    if "germany" in q_lower or "german" in q_lower:
        return "Germany"

    # LLM extraction fallback using Azure OpenAI if configured
    try:
        from langchain_openai import AzureChatOpenAI
        endpoint = os.getenv("AZURE_OPENAI_ENDPOINT")
        key = os.getenv("AZURE_OPENAI_API_KEY")
        deployment = os.getenv("AZURE_OPENAI_DEPLOYMENT", "gpt-5-mini-2")
        if endpoint and key:
            llm = AzureChatOpenAI(
                azure_endpoint=endpoint,
                api_key=key,
                azure_deployment=deployment,
                api_version=os.getenv("AZURE_OPENAI_API_VERSION", "v1"),
                temperature=0.0,
                max_tokens=30,
            )
            prompt = (
                f"Extract the primary location or country requested in this query. "
                f"If no specific location is mentioned, return 'India'. "
                f"Return ONLY the location string (e.g. 'India', 'United States'). Query: '{query}'"
            )
            res = await llm.ainvoke(prompt)
            loc = str(res.content).strip().strip('"').strip("'")
            if loc:
                return loc
    except Exception as exc:
        print(f"[WARN] Location extraction error: {exc}", flush=True)

    return "India"



# ── FastAPI App ────────────────────────────────────────────────────────────────

app = FastAPI(title="Agent Pipeline API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


visualizer_broker = VisualizerBroker()


# ── Helper Functions ───────────────────────────────────────────────────────────

async def _broadcast(state: AgentState, line: str) -> None:
    """Buffer line and push to all live SSE subscriber queues."""
    state.log_buffer.append(line)
    dead: List[asyncio.Queue] = []
    for q in list(state.log_queues):
        try:
            q.put_nowait(line)
        except asyncio.QueueFull:
            dead.append(q)
    for q in dead:
        if q in state.log_queues:
            state.log_queues.remove(q)

    # Also fan-out plain text logs to WS clients as a 'log' envelope
    ws_envelope = json.dumps({"event_type": "log", "data": {"message": line}})
    dead_ws = []
    for ws in list(_ws_log_clients):
        try:
            await ws.send_text(ws_envelope)
        except Exception:
            dead_ws.append(ws)
    for ws in dead_ws:
        if ws in _ws_log_clients:
            _ws_log_clients.remove(ws)


async def _send_sentinel(state: AgentState) -> None:
    """Signal all SSE queues that the stream has ended."""
    for q in list(state.log_queues):
        try:
            q.put_nowait(None)
        except Exception:
            pass


def _extract_url(cell_or_val) -> str:
    """Extract a plain URL from an openpyxl Cell, HYPERLINK formula, or URL string."""
    if cell_or_val is None:
        return ""

    # Check openpyxl Cell hyperlink attribute first
    if hasattr(cell_or_val, "hyperlink") and cell_or_val.hyperlink:
        target = getattr(cell_or_val.hyperlink, "target", "") or ""
        if target and ("http" in str(target).lower() or "linkedin.com" in str(target).lower()):
            val = str(target).strip()
            if not val.startswith("http://") and not val.startswith("https://"):
                val = "https://" + val.lstrip("/")
            return val

    # Get cell value if Cell object
    val_raw = cell_or_val.value if hasattr(cell_or_val, "value") else cell_or_val
    if not val_raw:
        return ""

    s = str(val_raw).strip()

    # Parse =HYPERLINK("url", "label") formula
    if s.startswith("="):
        m = re.search(r'HYPERLINK\(\s*["\']([^"\']+)["\']', s, re.IGNORECASE)
        if m:
            s = m.group(1).strip()

    if s.lower() in ("not available", "n/a", "not found", "none", "—", ""):
        return ""

    # Ensure http prefix if domain or linkedin URL
    if "linkedin.com" in s.lower() or "http" in s.lower() or s.startswith("www."):
        if not s.startswith("http://") and not s.startswith("https://"):
            s = "https://" + s.lstrip("/")
        return s

    return ""


def _extract_domain_from_url(url: str) -> str:
    """Extract bare domain (e.g. 'acmecorp.com') from a full URL or domain string."""
    if not url:
        return ""
    url = url.strip()
    try:
        from urllib.parse import urlparse
        parsed = urlparse(url if url.startswith("http") else f"https://{url}")
        host = parsed.hostname or ""
        return host.lstrip("www.")
    except Exception:
        return url.replace("https://", "").replace("http://", "").lstrip("www.").split("/")[0]


def _hunter_email_count(domain: str) -> int:
    """
    Calls Hunter.io Email Count API (GET /v2/email-count).
    Returns the total email count for the domain, or 0 on error / missing domain.
    This endpoint is FREE and does NOT consume Hunter credits.
    """
    if not domain:
        return 0
    try:
        import requests as _req
        key = get_hunter_api_key() if get_hunter_api_key else None
        if not key:
            return 0
        resp = _req.get(
            "https://api.hunter.io/v2/email-count",
            params={"domain": domain, "api_key": key},
            timeout=10,
        )
        data = resp.json()
        return int(data.get("data", {}).get("total", 0))
    except Exception as exc:
        print(f"[WARN] Hunter email-count error for '{domain}': {exc}", flush=True)
        return 0


def _parse_shortlister_excel(path: str) -> List[Dict]:
    """Read a Shortlister report Excel and return structured company dicts."""
    companies: List[Dict] = []
    try:
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [str(c.value) if c.value else "" for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=False):
            if not any(c.value for c in row):
                continue
            raw = {headers[i]: row[i].value for i in range(min(len(headers), len(row)))}
            companies.append({
                "company_name": str(raw.get("Company Name", "") or ""),
                "ticker":       str(raw.get("Ticker", "") or ""),
                "industry":     str(raw.get("Industry", "") or ""),
                "revenue":      str(raw.get("Revenue (TTM)", "") or ""),
                "net_income":   str(raw.get("Net Income (TTM)", "") or ""),
                "employees":    str(raw.get("Employee Count", "") or ""),
                "location":     str(raw.get("Location", "") or ""),
                "summary":      str(raw.get("Business Summary", "") or ""),
                "website":      _extract_url(raw.get("Website", "")),
                "linkedin_url": _extract_url(raw.get("LinkedIn (Verified)", "")),
                "linkedin_confirmed": str(raw.get("LinkedIn Confirmed?", "") or ""),
            })
    except Exception as exc:
        print(f"[WARN] Could not parse Excel at {path}: {exc}", flush=True)
    return companies


def _sync_all_outputs(agent_dir: Path, dest_dir: Path) -> List[str]:
    """
    Copy all generated .xlsx files (and .json files) from agent_dir/output to dest_dir.
    Ignores temporary Excel lock files starting with '~$'.
    """
    output_dir = agent_dir / "output"
    copied: List[str] = []
    if not output_dir.exists():
        return copied

    dest_dir.mkdir(parents=True, exist_ok=True)
    
    for f in output_dir.glob("*"):
        if f.name.startswith("~$") or f.is_dir() or f.name.endswith(".py"):
            continue
        if f.suffix in (".xlsx", ".json"):
            try:
                dest_file = dest_dir / f.name
                shutil.copy2(str(f), str(dest_file))
                copied.append(str(dest_file))
            except Exception as exc:
                print(f"[WARN] Could not copy {f.name} to {dest_dir}: {exc}", flush=True)

    return copied



async def _stream_subprocess(
    state: AgentState,
    cmd: List[str],
    cwd: Path,
    stdin_input: Optional[str] = None,
    prefix: Optional[str] = None,
    env_overrides: Optional[Dict[str, str]] = None,
) -> int:
    """
    Run a subprocess, tee its stdout+stderr to:
      1. The real terminal (unchanged, via print())
      2. All SSE subscriber queues (via _broadcast)

    Uses a reader thread + non-blocking queue poll to keep the asyncio
    event loop free on Windows where ProactorEventLoop subprocesses can
    behave inconsistently.
    """
    env = os.environ.copy()
    env["PYTHONUNBUFFERED"] = "1"  # force line-flushed output from subprocesses
    if env_overrides:
        env.update(env_overrides)

    line_q: "queue.Queue[Optional[str]]" = queue.Queue()

    proc = subprocess.Popen(
        cmd,
        stdin=subprocess.PIPE if stdin_input is not None else None,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        cwd=str(cwd),
        text=True,
        encoding="utf-8",
        errors="replace",
        env=env,
    )

    if stdin_input is not None:
        try:
            proc.stdin.write(stdin_input)
            proc.stdin.close()
        except Exception:
            pass

    def _reader():
        try:
            for raw_line in proc.stdout:
                line_q.put(raw_line.rstrip("\n"))
        finally:
            proc.wait()
            line_q.put(None)  # sentinel

    threading.Thread(target=_reader, daemon=True).start()

    while True:
        try:
            item = line_q.get_nowait()
        except queue.Empty:
            await asyncio.sleep(0.02)
            continue

        if item is None:
            break

        formatted_item = f"[{prefix}] {item}" if prefix else item

        # Tee to terminal
        print(formatted_item, flush=True)
        # Broadcast to browser
        await _broadcast(state, formatted_item)

    return proc.returncode


# ── Shortlister Background Task ────────────────────────────────────────────────

async def _run_shortlister_bg(query: str, run_id: str) -> None:
    state = shortlister_state
    state.status = "running"
    state.error = None
    state.user_query = query
    state.extracted_location = await _extract_location_from_query(query)
    await _broadcast(state, f"📍 Extracted target location from query: '{state.extracted_location}'")

    try:
        rc = await _stream_subprocess(
            state,
            [sys.executable, "-u", "main.py", query],
            SHORTLISTER_DIR,
            env_overrides={"PIPELINE_RUN_ID": run_id},
        )

        # Detect and sync Excel outputs to parent directory
        copied = _sync_all_outputs(
            SHORTLISTER_DIR, PARENT_OUTPUT_DIR / "shortlister"
        )
        _sync_all_outputs(
            LINKEDIN_DIR, PARENT_OUTPUT_DIR / "linkedin"
        )
        if copied:
            consolidated = [p for p in copied if "consolidated" in Path(p).name.lower()]
            newest = max(consolidated if consolidated else copied, key=lambda p: os.path.getmtime(p))
            state.last_excel = newest
            state.companies = _parse_shortlister_excel(newest)

        state.status = "done" if rc == 0 else "error"
        if rc != 0:
            state.error = f"Process exited with code {rc}"

    except Exception as exc:
        state.status = "error"
        state.error = str(exc)
        await _broadcast(state, f"[ERROR] {exc}")
    finally:
        await _send_sentinel(state)
        await visualizer_broker.mark_process_done(run_id, state.status, state.error)


# ── LinkedIn Finder Background Task ───────────────────────────────────────────

async def _run_linkedin_bg(companies: list, concurrency: int = 10) -> None:
    state = linkedin_state
    state.status = "running"
    state.error = None
    state.total_count = len(companies)
    state.completed_count = 0
    state.active_companies = []

    sem = asyncio.Semaphore(max(1, concurrency))
    lock = asyncio.Lock()
    failed_companies: List[str] = []

    await _broadcast(state, f"🚀 Starting parallel extraction for {len(companies)} companies with concurrency limit = {concurrency}")

    async def _worker(company: dict):
        company_name = company["company_name"]
        linkedin_url = company.get("linkedin_url", "")
        website = company.get("website", "")
        target_location = company.get("location") or shortlister_state.extracted_location or "India"

        async with sem:
            async with lock:
                state.active_companies.append(company_name)
                state.current_company = ", ".join(state.active_companies)

            separator = f"{'─' * 60}"
            await _broadcast(state, f"[{company_name}] {separator}")
            await _broadcast(state, f"[{company_name}]   Processing: {company_name} (Location: {target_location})")
            await _broadcast(state, f"[{company_name}] {separator}")

            cmd = [
                sys.executable, "-u", "main.py",
                "--company_name", company_name,
                "--company_website", website or "",
                "--location", target_location or "India",
            ]
            if linkedin_url and linkedin_url != "NOT FOUND":
                cmd.extend(["--company_linkedin", linkedin_url])

            try:
                rc = await _stream_subprocess(
                    state,
                    cmd,
                    LINKEDIN_DIR,
                    prefix=company_name,
                )
                
                async with lock:
                    _sync_all_outputs(LINKEDIN_DIR, PARENT_OUTPUT_DIR / "linkedin")

                if rc != 0:
                    async with lock:
                        failed_companies.append(f"'{company_name}' (exit code {rc})")
                    await _broadcast(state, f"[{company_name}] ❌ Exited with code {rc}")
                else:
                    await _broadcast(state, f"[{company_name}] ✅ Completed successfully")

            except Exception as exc:
                async with lock:
                    failed_companies.append(f"'{company_name}' ({exc})")
                await _broadcast(state, f"[{company_name}] ❌ Error: {exc}")
            finally:
                async with lock:
                    if company_name in state.active_companies:
                        state.active_companies.remove(company_name)
                    state.completed_count += 1
                    state.current_company = ", ".join(state.active_companies) if state.active_companies else None

    tasks = [_worker(comp) for comp in companies]
    await asyncio.gather(*tasks, return_exceptions=True)

    # Sync all outputs at end of batch
    _sync_all_outputs(LINKEDIN_DIR, PARENT_OUTPUT_DIR / "linkedin")

    if failed_companies:
        state.status = "error" if len(failed_companies) == len(companies) else "done"
        state.error = f"Completed with errors in {len(failed_companies)} company(ies): {', '.join(failed_companies)}"
    else:
        state.status = "done"

    state.active_companies.clear()
    state.current_company = None
    await _send_sentinel(state)


# ── Generic SSE Generator ──────────────────────────────────────────────────────

def _make_sse_generator(state: AgentState, request: Request):
    async def generator():
        # Replay buffered lines for reconnecting clients
        for line in list(state.log_buffer):
            yield {"event": "log", "data": line}

        # If already finished, send done and exit immediately
        if state.status not in ("idle", "running"):
            yield {"event": "done", "data": state.status}
            return

        # If still idle (run not started), nothing to stream
        if state.status == "idle":
            return

        # Subscribe to live updates
        q: asyncio.Queue = asyncio.Queue(maxsize=2000)
        state.log_queues.append(q)
        try:
            while True:
                if await request.is_disconnected():
                    break
                try:
                    item = await asyncio.wait_for(q.get(), timeout=15.0)
                    if item is None:
                        yield {"event": "done", "data": state.status}
                        break
                    yield {"event": "log", "data": item}
                except asyncio.TimeoutError:
                    yield {"event": "ping", "data": ""}
        finally:
            if q in state.log_queues:
                state.log_queues.remove(q)

    return generator


# ── REST Models ────────────────────────────────────────────────────────────────

class ShortlisterRequest(BaseModel):
    query: str


class LinkedInCompany(BaseModel):
    company_name: str
    linkedin_url: str
    website: str


class LinkedInRequest(BaseModel):
    companies: List[LinkedInCompany]
    concurrency: int = 10


# ── WebSocket structured log streaming endpoint ───────────────────────────────

@app.websocket("/ws/visualizer")
async def ws_visualizer(websocket: WebSocket):
    """Stream live Shortlister JSONL telemetry to every connected browser."""
    await websocket.accept()
    client_queue = visualizer_broker.subscribe()
    try:
        while True:
            try:
                payload = await asyncio.wait_for(client_queue.get(), timeout=25.0)
                await websocket.send_json(payload)
            except asyncio.TimeoutError:
                await websocket.send_json({
                    "stream_version": "1.0",
                    "type": "keepalive",
                    "timestamp": datetime.now(timezone.utc).isoformat(),
                })
    except (WebSocketDisconnect, Exception):
        pass
    finally:
        visualizer_broker.unsubscribe(client_queue)


@app.websocket("/ws/logs")
async def ws_logs(websocket: WebSocket):
    """
    WebSocket endpoint that streams structured JSON log events emitted by
    the pipeline agents in real-time.

    Events are JSON objects matching the envelope schema:
    {
        "schema_version": "1.0",
        "timestamp": "<ISO-8601>",
        "level": "INFO",
        "run_id": "<session-id>",
        "agent": "shortlister" | "linkedin",
        "component": "<module>",
        "event_type": "pipeline_start" | "node_enter" | "node_exit" | "tool_start" |
                      "tool_end" | "tool_error" | "llm_start" | "llm_end" |
                      "chain_start" | "chain_end" | "pipeline_end" | "pipeline_error" |
                      "progress" | "company_enriched" | "company_dropped" |
                      "contact_found" | "company_contacts_snapshot" | "subagent_start" |
                      "subagent_error" | "subagent_end" | "search_query" | "mcp_connect" |
                      "mcp_error" | "log",
        "correlation_id": null,
        "data": { ... event-specific payload ... }
    }
    """
    await websocket.accept()
    _ws_log_clients.append(websocket)
    try:
        # Drain the shared ws_queue and forward each JSON line to this client.
        # Also relay any plain-text logs that arrive from _broadcast().
        while True:
            try:
                item = await asyncio.wait_for(_ws_log_queue.get(), timeout=30.0)
                if item is None:
                    await websocket.send_text('{"event_type":"stream_end"}')
                    break
                await websocket.send_text(item)
            except asyncio.TimeoutError:
                # Keepalive ping
                try:
                    await websocket.send_text('{"event_type":"keepalive"}')
                except Exception:
                    break
    except (WebSocketDisconnect, Exception):
        pass
    finally:
        if websocket in _ws_log_clients:
            _ws_log_clients.remove(websocket)


# ── Shortlister Endpoints ──────────────────────────────────────────────────────

@app.post("/shortlister/run")
async def run_shortlister(req: ShortlisterRequest, background_tasks: BackgroundTasks):
    if shortlister_state.status == "running":
        raise HTTPException(409, "Shortlister is already running")
    shortlister_state.reset()
    run_id = f"session-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:6]}"
    log_path = SHORTLISTER_DIR / "logs" / f"run_{run_id}.jsonl"
    await visualizer_broker.begin_run(run_id, log_path)
    background_tasks.add_task(_run_shortlister_bg, req.query, run_id)
    return {"status": "started", "run_id": run_id}


@app.get("/shortlister/stream")
async def shortlister_stream(request: Request):
    return EventSourceResponse(_make_sse_generator(shortlister_state, request)())


@app.get("/shortlister/status")
async def shortlister_status():
    return {
        "status": shortlister_state.status,
        "excel_path": shortlister_state.last_excel,
        "companies": shortlister_state.companies,
        "error": shortlister_state.error,
    }


@app.get("/shortlister/download")
async def shortlister_download():
    path = shortlister_state.last_excel
    if not path or not Path(path).exists():
        raise HTTPException(404, "No Excel file available yet")
    p = Path(path)
    return FileResponse(
        str(p),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=p.name,
    )


# ── LinkedIn Finder Endpoints ──────────────────────────────────────────────────

@app.post("/linkedin/run")
async def run_linkedin(req: LinkedInRequest, background_tasks: BackgroundTasks):
    if linkedin_state.status == "running":
        raise HTTPException(409, "LinkedIn Finder is already running")
    linkedin_state.reset()
    companies_dicts = [c.model_dump() for c in req.companies]
    background_tasks.add_task(_run_linkedin_bg, companies_dicts, req.concurrency)
    return {"status": "started", "concurrency": req.concurrency}


@app.get("/linkedin/stream")
async def linkedin_stream(request: Request):
    return EventSourceResponse(_make_sse_generator(linkedin_state, request)())


@app.get("/linkedin/status")
async def linkedin_status():
    return {
        "status": linkedin_state.status,
        "current_company": linkedin_state.current_company,
        "active_companies": linkedin_state.active_companies,
        "completed_count": linkedin_state.completed_count,
        "total_count": linkedin_state.total_count,
        "error": linkedin_state.error,
    }


@app.get("/linkedin/files")
async def linkedin_files():
    """List all available company Excel files in output/linkedin."""
    output_dir = PARENT_OUTPUT_DIR / "linkedin"
    if not output_dir.exists():
        return {"files": []}

    files = []
    for f in sorted(output_dir.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True):
        if f.name.startswith("~$"):
            continue
        files.append({
            "filename": f.name,
            "size_bytes": f.stat().st_size,
            "mtime": f.stat().st_mtime,
            "download_url": f"/linkedin/download/{f.name}"
        })
    return {"files": files}


@app.get("/linkedin/download/{filename}")
async def linkedin_download_file(filename: str):
    """Download a specific company Excel report from output/linkedin."""
    safe_name = Path(filename).name
    file_path = PARENT_OUTPUT_DIR / "linkedin" / safe_name
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(404, f"File '{safe_name}' not found")
    return FileResponse(
        str(file_path),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=safe_name,
    )


# ── Hunter & Decision Makers Endpoints ─────────────────────────────────────────

def _get_current_session_consolidated_excel() -> Optional[Path]:
    """Find the consolidated Excel report generated in the active/latest session."""
    # 1. Check if shortlister_state has last_excel recorded
    if shortlister_state.last_excel:
        p = Path(shortlister_state.last_excel)
        if p.exists() and "consolidated" in p.name.lower():
            return p
        # If last_excel is non-consolidated, search for matching consolidated file in output/shortlister
        shortlister_dir = PARENT_OUTPUT_DIR / "shortlister"
        if shortlister_dir.exists():
            matched = [f for f in shortlister_dir.glob("report_consolidated_*.xlsx") if not f.name.startswith("~$")]
            if matched:
                return max(matched, key=lambda f: f.stat().st_mtime)

    # 2. Fallback: find the single newest report_consolidated_*.xlsx in output/shortlister or output/
    shortlister_dir = PARENT_OUTPUT_DIR / "shortlister"
    candidates: List[Path] = []
    if shortlister_dir.exists():
        candidates.extend(list(shortlister_dir.glob("report_consolidated_*.xlsx")))
    candidates.extend(list(PARENT_OUTPUT_DIR.glob("report_consolidated_*.xlsx")))

    valid_candidates = [f for f in candidates if not f.name.startswith("~$")]
    if valid_candidates:
        return max(valid_candidates, key=lambda f: f.stat().st_mtime)

    return None


def _get_all_decision_makers() -> List[Dict]:
    """
    Parse decision makers strictly from the Consolidated Excel file generated
    in the active/latest session. Supports both collapsible grouped layout and
    legacy flat layout. Cross-references with LinkedIn Agent outputs.
    """
    latest_excel = _get_current_session_consolidated_excel()
    dm_map: Dict[tuple, Dict] = {}

    if latest_excel and latest_excel.exists():
        try:
            wb = openpyxl.load_workbook(str(latest_excel), data_only=False)
            if "Decision Maker Contacts" in wb.sheetnames:
                ws = wb["Decision Maker Contacts"]
                
                # Determine format: Check if row 1 is flat header or collapsible header
                is_flat_header = False
                if ws.max_row >= 1:
                    row1_vals = [str(c.value or "").strip() for c in ws[1]]
                    if "Company Name" in row1_vals and "Decision Maker Name" in row1_vals:
                        is_flat_header = True

                if is_flat_header:
                    headers = [str(c.value or "").strip() for c in ws[1]]
                    col_map = {h: i for i, h in enumerate(headers)}
                    for row in ws.iter_rows(min_row=2, values_only=False):
                        if not any(c.value for c in row):
                            continue
                        c_n = str(row[col_map.get("Company Name", 0)].value or "").strip() if "Company Name" in col_map else ""
                        c_w = _extract_url(row[col_map.get("Company Website", 1)].value) if "Company Website" in col_map else ""
                        c_l = _extract_url(row[col_map.get("Company LinkedIn", 2)].value) if "Company LinkedIn" in col_map else ""
                        l_c = str(row[col_map.get("Location", 3)].value or "").strip() if "Location" in col_map else ""
                        p_s = str(row[col_map.get("Decision Maker Position", 4)].value or "").strip() if "Decision Maker Position" in col_map else ""
                        d_n = str(row[col_map.get("Decision Maker Name", 5)].value or "").strip() if "Decision Maker Name" in col_map else ""
                        d_l = _extract_url(row[col_map.get("Decision Maker LinkedIn", 6)].value) if "Decision Maker LinkedIn" in col_map else ""

                        if d_n and d_n not in ("—", "N/A", "None", ""):
                            key = (c_n.lower(), d_n.lower())
                            if key not in dm_map:
                                dm_map[key] = {
                                    "company_name": c_n,
                                    "company_website": c_w,
                                    "company_linkedin": c_l,
                                    "location": l_c,
                                    "name": d_n,
                                    "position": p_s,
                                    "linkedin_url": d_l,
                                    "source_file": latest_excel.name
                                }
                else:
                    # Collapsible grouped layout parser
                    current_company = ""
                    for row in ws.iter_rows(min_row=1, values_only=False):
                        row_cells = [c for c in row]
                        row_vals = [str(c.value or "").strip() for c in row_cells]
                        if not any(row_vals):
                            continue

                        val0 = row_vals[0]

                        # 1. Company Banner Row (e.g. "▸ Reliance Industries  (5 contacts)")
                        if val0.startswith("▸") or "contacts)" in val0.lower():
                            m = re.search(r'▸\s*(.*?)\s*\(\d+\s*contacts\)', val0, re.IGNORECASE)
                            if m:
                                current_company = m.group(1).strip()
                            else:
                                current_company = val0.lstrip("▸").strip().split("(")[0].strip()
                            continue

                        # 2. Sub-header row
                        if val0 in ("#", "Company Name") or "LinkedIn" in str(row_vals[3] if len(row_vals) > 3 else ""):
                            continue

                        # 3. Data row: [# (digit), Name, Position, LinkedIn (Person), Company LinkedIn, Location]
                        if val0.isdigit() or (isinstance(row_cells[0].value, (int, float)) and int(row_cells[0].value) > 0):
                            dm_name = row_vals[1] if len(row_vals) > 1 else ""
                            pos = row_vals[2] if len(row_vals) > 2 else ""
                            dm_li = _extract_url(row_cells[3]) if len(row_cells) > 3 else ""
                            co_li = _extract_url(row_cells[4]) if len(row_cells) > 4 else ""
                            loc = row_vals[5] if len(row_vals) > 5 else ""

                            if dm_name and dm_name not in ("—", "N/A", "None", ""):
                                key = (current_company.lower(), dm_name.lower())
                                if key not in dm_map:
                                    dm_map[key] = {
                                        "company_name": current_company,
                                        "company_website": "",
                                        "company_linkedin": co_li,
                                        "location": loc,
                                        "name": dm_name,
                                        "position": pos,
                                        "linkedin_url": dm_li,
                                        "source_file": latest_excel.name
                                    }
        except Exception as exc:
            print(f"[WARN] Error reading session Consolidated Excel {latest_excel.name}: {exc}", flush=True)

    # JSON fallback if dm_map is empty
    if not dm_map:
        shortlister_dir = PARENT_OUTPUT_DIR / "shortlister"
        json_candidates = []
        if shortlister_dir.exists():
            json_candidates.extend(list(shortlister_dir.glob("report_*.json")))
        json_candidates.extend(list(PARENT_OUTPUT_DIR.glob("report_*.json")))
        if json_candidates:
            latest_json = max(json_candidates, key=lambda f: f.stat().st_mtime)
            try:
                with open(latest_json, "r", encoding="utf-8") as f:
                    jdata = json.load(f)
                    for comp in jdata.get("companies", []):
                        c_n = comp.get("name", "") or comp.get("company_name", "")
                        c_w = comp.get("website", "") or comp.get("company_website", "")
                        c_l = comp.get("linkedin_verified", "") or comp.get("company_linkedin", "")
                        l_c = comp.get("location", "")
                        for dm in comp.get("decision_makers", []):
                            d_n = dm.get("name", "")
                            p_s = dm.get("position") or dm.get("title") or ""
                            d_l = dm.get("linkedin_url") or dm.get("profile_url") or ""
                            if d_n and d_n not in ("—", "N/A", "None", ""):
                                k = (c_n.lower(), d_n.lower())
                                if k not in dm_map:
                                    dm_map[k] = {
                                        "company_name": c_n,
                                        "company_website": c_w,
                                        "company_linkedin": c_l,
                                        "location": l_c,
                                        "name": d_n,
                                        "position": p_s,
                                        "linkedin_url": d_l,
                                        "source_file": latest_json.name
                                    }
            except Exception as exc:
                print(f"[WARN] Error reading JSON report {latest_json.name}: {exc}", flush=True)

    # Cross-reference with output/linkedin to populate any missing LinkedIn URLs
    linkedin_dir = PARENT_OUTPUT_DIR / "linkedin"
    if linkedin_dir.exists():
        for ef in linkedin_dir.glob("contacts_*.xlsx"):
            if ef.name.startswith("~$"):
                continue
            try:
                wb_link = openpyxl.load_workbook(str(ef), data_only=False)
                ws_link = wb_link.active
                headers_link = [str(c.value or "").strip() for c in ws_link[1]]
                col_map_link = {h: i for i, h in enumerate(headers_link)}
                for row in ws_link.iter_rows(min_row=2, values_only=False):
                    c_n = str(row[col_map_link.get("Company Name", 0)].value or "").strip() if "Company Name" in col_map_link else ""
                    d_n = str(row[col_map_link.get("Decision Maker Name", 5)].value or "").strip() if "Decision Maker Name" in col_map_link else ""
                    d_li = _extract_url(row[col_map_link.get("Decision Maker LinkedIn", 6)].value) if "Decision Maker LinkedIn" in col_map_link else ""
                    key = (c_n.lower(), d_n.lower())
                    if key in dm_map and not dm_map[key]["linkedin_url"] and d_li:
                        dm_map[key]["linkedin_url"] = d_li
            except Exception:
                pass

    dms: List[Dict] = []
    for idx, item in enumerate(dm_map.values(), 1):
        item["id"] = f"dm-{idx}"
        dms.append(item)

    return dms


@app.get("/linkedin/decision-makers")
async def get_decision_makers():
    """Retrieve all parsed decision makers across all companies processed so far."""
    dms = _get_all_decision_makers()
    return {"decision_makers": dms, "total": len(dms)}


class HunterEmailRequest(BaseModel):
    linkedin_url: str
    first_name: Optional[str] = None
    last_name: Optional[str] = None
    full_name: Optional[str] = None
    domain: Optional[str] = None
    company: Optional[str] = None


class StoreEmailRequest(BaseModel):
    company_name: str
    dm_name: str
    email: str


class HunterSmartEnrichDM(BaseModel):
    id: str
    name: Optional[str] = None
    linkedin_url: Optional[str] = None
    company_name: Optional[str] = None
    company_website: Optional[str] = None
    position: Optional[str] = None


class HunterSmartEnrichRequest(BaseModel):
    decision_makers: List[HunterSmartEnrichDM]


@app.post("/hunter/find-email")
async def hunter_find_email(req: HunterEmailRequest):
    """Retrieve email address via Hunter.io API for a specific Decision Maker's LinkedIn profile/handle."""
    if find_email_by_linkedin is None:
        raise HTTPException(500, "Hunter LinkedIn Email Finder module not loaded")

    if not req.linkedin_url or not req.linkedin_url.strip():
        raise HTTPException(400, "LinkedIn URL or handle is required")

    try:
        res = await asyncio.to_thread(
            find_email_by_linkedin,
            linkedin_input=req.linkedin_url.strip(),
            first_name=req.first_name,
            last_name=req.last_name,
            full_name=req.full_name,
            domain=req.domain,
            company=req.company,
        )
        return res
    except ValueError as val_err:
        raise HTTPException(400, str(val_err))
    except Exception as exc:
        raise HTTPException(500, f"Hunter API error: {exc}")


@app.get("/hunter/email-count")
async def hunter_email_count(domain: str):
    """
    Check how many emails Hunter.io has indexed for a given domain.
    FREE endpoint — does not consume Hunter credits.
    Returns { domain, count }.
    """
    count = await asyncio.to_thread(_hunter_email_count, domain)
    return {"domain": domain, "count": count}


@app.post("/hunter/smart-enrich")
async def hunter_smart_enrich(req: HunterSmartEnrichRequest):
    """
    Smart sequential email enrichment:
    1. Groups DMs by company domain.
    2. Calls Email Count per domain — skips companies where count == 0.
    3. For each eligible company, sequentially calls Email Finder on DMs (with LinkedIn URLs).
    4. Stops per company once 3 emails have been found.
    5. Returns enriched DMs pinned first, metadata about skipped/stopped companies.
    """
    if find_email_by_linkedin is None:
        raise HTTPException(500, "Hunter LinkedIn Email Finder module not loaded")

    EMAIL_CAP = 3  # hardcoded per-company cap

    # ── Group DMs by domain ─────────────────────────────────────────────────────
    # domain → list of DMs (preserving original order)
    domain_groups: Dict[str, List[HunterSmartEnrichDM]] = {}
    dm_domains: Dict[str, str] = {}  # dm.id → domain

    for dm in req.decision_makers:
        domain = _extract_domain_from_url(dm.company_website or "")
        dm_domains[dm.id] = domain
        if domain not in domain_groups:
            domain_groups[domain] = []
        domain_groups[domain].append(dm)

    enriched: List[Dict] = []          # DMs with found emails
    skipped_companies: List[str] = []  # count == 0
    stopped_at_cap: List[str] = []    # hit the 3-email cap
    no_indexed: set = set()            # dm IDs whose company had count=0

    # ── Process each domain group ───────────────────────────────────────────────
    for domain, dms in domain_groups.items():
        company_label = dms[0].company_name or domain or "Unknown"

        # Step 1: Email Count gate
        count = await asyncio.to_thread(_hunter_email_count, domain)
        if count == 0:
            print(f"[SMART-ENRICH] ⛔ {company_label} — email count=0, skipping", flush=True)
            skipped_companies.append(company_label)
            for dm in dms:
                no_indexed.add(dm.id)
            continue

        print(f"[SMART-ENRICH] ✅ {company_label} — email count={count}, proceeding", flush=True)

        # Step 2: Sequential Email Finder, stop at EMAIL_CAP
        found = 0
        for dm in dms:
            if found >= EMAIL_CAP:
                break
            if not dm.linkedin_url or not dm.linkedin_url.strip():
                continue  # skip DMs without LinkedIn URL silently

            try:
                result = await asyncio.to_thread(
                    find_email_by_linkedin,
                    linkedin_input=dm.linkedin_url.strip(),
                    full_name=dm.name,
                    domain=domain,
                    company=dm.company_name,
                )
                email = result.get("data", {}).get("email") if isinstance(result.get("data"), dict) else None
                score = result.get("data", {}).get("score") if isinstance(result.get("data"), dict) else None

                if email:
                    print(f"[SMART-ENRICH]   📧 Found email for {dm.name}: {email}", flush=True)
                    enriched.append({
                        "id": dm.id,
                        "email": email,
                        "score": score,
                        "raw": result,
                    })
                    found += 1
                else:
                    print(f"[SMART-ENRICH]   — No email for {dm.name}", flush=True)

            except Exception as exc:
                print(f"[SMART-ENRICH]   ⚠️ Error for {dm.name}: {exc}", flush=True)

        if found >= EMAIL_CAP:
            stopped_at_cap.append(company_label)

    enriched_ids = {e["id"] for e in enriched}

    return {
        "enriched": enriched,
        "no_indexed_dm_ids": list(no_indexed),
        "skipped_companies": skipped_companies,
        "stopped_at_cap_companies": stopped_at_cap,
        "total_emails_found": len(enriched_ids),
    }


@app.post("/hunter/store-email")
async def store_email(req: StoreEmailRequest):
    """
    Persist a Hunter-retrieved email to the in-memory store so it is
    available for the consolidated Excel export without re-fetching.
    Called by the frontend immediately after every successful email find.
    """
    if not req.email or not req.email.strip():
        raise HTTPException(400, "email must not be empty")
    key = (req.company_name.strip().lower(), req.dm_name.strip().lower())
    _email_store[key] = req.email.strip()
    return {"stored": True, "key": f"{req.company_name} / {req.dm_name}"}


@app.get("/export/consolidated")
async def export_consolidated():
    """
    Generate and download the 2-sheet consolidated Excel report:
      Sheet 1 — Company Overview
      Sheet 2 — Decision Makers (PivotTable or grouped table)
    Also saves a copy to output/exports/.
    """
    if build_export_xlsx is None:
        raise HTTPException(500, "export_builder module not loaded")

    companies = shortlister_state.companies or []
    location  = shortlister_state.extracted_location or "India"
    dms       = _get_all_decision_makers()

    try:
        xlsx_bytes = await asyncio.to_thread(
            build_export_xlsx, companies, location, dms, dict(_email_store)
        )
    except Exception as exc:
        raise HTTPException(500, f"Export generation failed: {exc}")

    # Save a copy to disk
    from datetime import datetime as _dt
    timestamp  = _dt.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename   = f"export_{timestamp}.xlsx"
    disk_path  = PARENT_OUTPUT_DIR / "exports" / filename
    try:
        disk_path.write_bytes(xlsx_bytes)
        print(f"[EXPORT] Saved to {disk_path}", flush=True)
    except Exception as save_exc:
        print(f"[WARN] Could not save export to disk: {save_exc}", flush=True)

    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/export/list")
async def list_exports():
    """List all previously exported Excel files in output/exports/."""
    exports_dir = PARENT_OUTPUT_DIR / "exports"
    if not exports_dir.exists():
        return {"files": []}
    files = [
        {
            "filename": f.name,
            "size_bytes": f.stat().st_size,
            "mtime": f.stat().st_mtime,
            "download_url": f"/export/download/{f.name}",
        }
        for f in sorted(exports_dir.glob("export_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not f.name.startswith("~$")
    ]
    return {"files": files}


@app.get("/export/download/{filename}")
async def download_export(filename: str):
    """Re-download a previously generated export file."""
    safe = Path(filename).name
    path = PARENT_OUTPUT_DIR / "exports" / safe
    if not path.exists():
        raise HTTPException(404, f"Export '{safe}' not found")
    return FileResponse(
        str(path),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=safe,
    )


# ── Utility Endpoints ──────────────────────────────────────────────────────────

@app.post("/reset")
async def reset_all():
    """Reset both agent states (does not kill running processes)."""
    shortlister_state.reset()
    linkedin_state.reset()
    return {"status": "reset"}


@app.get("/health")
async def health():
    return {"status": "ok"}
