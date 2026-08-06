import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
import config

def export_contacts_excel(
    company_name: str,
    company_website: str,
    company_linkedin: str,
    location: str,
    decision_makers: list,
    session_id: str
) -> str:
    """
    Exports decision maker contacts to an Excel file with the specified 7 columns:
    1. Company Name
    2. Company Website
    3. Company LinkedIn
    4. Location
    5. Decision Maker Position
    6. Decision Maker Name
    7. Decision Maker LinkedIn
    """
    os.makedirs(config.OUTPUT_DIR, exist_ok=True)
    safe_company = "".join(c for c in company_name if c.isalnum() or c in (" ", "-", "_")).strip().replace(" ", "_")
    if not safe_company:
        safe_company = "Company"
        
    if safe_company.lower() in session_id.lower():
        filename = f"contacts_{session_id}.xlsx"
    else:
        filename = f"contacts_{safe_company}_{session_id}.xlsx"
        
    filepath = os.path.join(config.OUTPUT_DIR, filename)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Decision Makers"
    
    headers = [
        "Company Name",
        "Company Website",
        "Company LinkedIn",
        "Location",
        "Decision Maker Position",
        "Decision Maker Name",
        "Decision Maker LinkedIn"
    ]
    ws.append(headers)
    
    # Header styling
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for dm in decision_makers:
        pos = dm.get("position", "") or "Unknown"
        name = dm.get("name", "")
        dm_linkedin = dm.get("linkedin_url", "")
        
        row = [
            company_name,
            company_website,
            company_linkedin,
            location,
            pos,
            name,
            dm_linkedin
        ]
        ws.append(row)
        curr_row = ws.max_row
        
        # Format Hyperlinks
        if company_website:
            cell_web = ws.cell(row=curr_row, column=2)
            cell_web.value = f'=HYPERLINK("{company_website}", "{company_website}")'
            cell_web.font = Font(color="0563C1", underline="single")
            
        if company_linkedin:
            cell_comp_li = ws.cell(row=curr_row, column=3)
            cell_comp_li.value = f'=HYPERLINK("{company_linkedin}", "{company_linkedin}")'
            cell_comp_li.font = Font(color="0563C1", underline="single")
            
        if dm_linkedin:
            cell_dm_li = ws.cell(row=curr_row, column=7)
            cell_dm_li.value = f'=HYPERLINK("{dm_linkedin}", "{dm_linkedin}")'
            cell_dm_li.font = Font(color="0563C1", underline="single")

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    wb.save(filepath)
    return filepath
