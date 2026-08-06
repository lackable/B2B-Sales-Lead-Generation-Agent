"""
INPUT:  Excel with "Decision Maker Contacts" sheet (flat table, contacts repeated per company)
OUTPUT: Excel with AutoFilter on Company Name → filter dropdown in Excel to pick company
"""
 
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
 
# ── CONFIG ──────────────────────────────────────────────────
INPUT_FILE  = "C:\\Users\\saksh\\OneDrive\\Desktop\\AI\\Testing\\Shortlister Agent\\output\\report_consolidated_2026-08-04_14-27-36.xlsx"
SHEET_NAME  = "Decision Maker Contacts"
OUTPUT_FILE = "C:\\Users\\saksh\\OneDrive\\Desktop\\AI\\Testing\\Shortlister Agent\\output\\Decision_Maker_Pivot_Filtered.xlsx"
 
# Columns to keep in pivot (change order here if needed)
PIVOT_COLS = [
    "Company Name",
    "Decision Maker Name",
    "Decision Maker Position",
    "Decision Maker LinkedIn",
    "Location",
]
 
# ── READ ────────────────────────────────────────────────────
df = pd.read_excel(INPUT_FILE, sheet_name=SHEET_NAME)
df = df[PIVOT_COLS].fillna("—")
df = df.sort_values("Company Name").reset_index(drop=True)
 
# ── STYLES ──────────────────────────────────────────────────
FONT        = "Arial"
header_font = Font(name=FONT, bold=True, size=11, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="2F5496")
data_font   = Font(name=FONT, size=10)
border      = Border(
    left=Side("thin", color="D9D9D9"), right=Side("thin", color="D9D9D9"),
    top=Side("thin", color="D9D9D9"),  bottom=Side("thin", color="D9D9D9"),
)
alt_fill    = PatternFill("solid", fgColor="F2F2F2")  # alternating row color
center      = Alignment(horizontal="center", vertical="center")
wrap        = Alignment(vertical="center", wrap_text=True)
 
# ── BUILD WORKBOOK ──────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Pivot - Filter by Company"
 
# Header row
for ci, col_name in enumerate(PIVOT_COLS, 1):
    cell = ws.cell(row=1, column=ci, value=col_name)
    cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border
 
# Data rows
for ri, (_, row_data) in enumerate(df.iterrows(), 2):
    for ci, col_name in enumerate(PIVOT_COLS, 1):
        cell = ws.cell(row=ri, column=ci, value=row_data[col_name])
        cell.font, cell.border = data_font, border
        cell.alignment = center if ci == 1 else wrap
        if ri % 2 == 0:
            cell.fill = alt_fill
 
# Column widths
widths = {"A": 35, "B": 30, "C": 55, "D": 45, "E": 12}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
 
# ── AUTOFILTER (the actual filter dropdown) ─────────────────
last_row = ws.max_row
last_col = get_column_letter(len(PIVOT_COLS))
ws.auto_filter.ref = f"A1:{last_col}{last_row}"
 
# Freeze header row → stays visible when scrolling
ws.freeze_panes = "A2"
ws.sheet_view.showGridLines = False
 
# ── SAVE ────────────────────────────────────────────────────
wb.save(OUTPUT_FILE)
print(f"✓ Saved → {OUTPUT_FILE}")
print(f"  {len(df)} contacts | {df['Company Name'].nunique()} companies")
print(f"  Open in Excel → click Company Name dropdown → pick company")