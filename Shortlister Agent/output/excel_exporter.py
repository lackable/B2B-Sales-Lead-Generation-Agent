import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
import config

def export_excel(companies: list, timestamp: str) -> str:
    """Exports shortlisted and verified companies into a clean Excel spreadsheet."""
    os.makedirs(config.EXCEL_OUTPUT_DIR, exist_ok=True)
    filename = f"report_{timestamp}.xlsx"
    filepath = os.path.join(config.EXCEL_OUTPUT_DIR, filename)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Verified Shortlist"
    
    headers = [
        "Company Name", "Ticker", "Industry", "Revenue (TTM)", 
        "Net Income (TTM)", "Employee Count", "Location", "Business Summary", 
        "Website", "LinkedIn (Verified)", "LinkedIn Confirmed?"
    ]
    ws.append(headers)
    
    # Style headers
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for company in companies:
        is_confirmed = company.get("linkedin_confirmed")
        confirmed_str = "Verified" if is_confirmed else "Non-Verified"
        
        row = [
            company.get("name", "") or company.get("company_name", ""),
            company.get("ticker", ""),
            company.get("industry", "") or company.get("field", ""),
            company.get("revenue_ttm", "") or company.get("revenue", ""),
            company.get("net_income_ttm", "") or company.get("profit", ""),
            company.get("employee_count", "") or company.get("num_employees", ""),
            company.get("location", ""),
            company.get("business_summary", "") or company.get("summary", ""),
            company.get("website", "") or company.get("company_website", ""),
            company.get("linkedin_verified", "") or company.get("company_linkedin", ""),
            confirmed_str
        ]
        
        ws.append(row)
        curr_row = ws.max_row
        
        # Format Website hyperlink (Col 9)
        website = company.get("website", "") or company.get("company_website", "")
        if website and str(website).startswith("http"):
            ws.cell(row=curr_row, column=9).value = f'=HYPERLINK("{website}", "{website}")'
            ws.cell(row=curr_row, column=9).font = Font(color="0563C1", underline="single")
            
        # Format Verified LinkedIn hyperlink (Col 10)
        linkedin_ver = company.get("linkedin_verified", "") or company.get("company_linkedin", "")
        if linkedin_ver and str(linkedin_ver).startswith("http"):
            ws.cell(row=curr_row, column=10).value = f'=HYPERLINK("{linkedin_ver}", "{linkedin_ver}")'
            ws.cell(row=curr_row, column=10).font = Font(color="0563C1", underline="single")

        # Format LinkedIn Confirmed status cell (Col 11)
        conf_cell = ws.cell(row=curr_row, column=11)
        if is_confirmed:
            conf_cell.font = Font(color="0F5132", bold=True)
            conf_cell.fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
        else:
            conf_cell.font = Font(color="842029", bold=True)
            conf_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        conf_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                val_str = str(cell.value or "")
                if len(val_str) > max_length and not val_str.startswith("="):
                    max_length = len(val_str)
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_length + 3, 12), 60)

    wb.save(filepath)
    return filepath
