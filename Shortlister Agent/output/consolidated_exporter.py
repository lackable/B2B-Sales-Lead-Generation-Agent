"""
Consolidated Exporter
=====================
Builds 2-sheet consolidated Excel workbook:
  Sheet 1 — "Verified Shortlist"      : Top shortlisted companies with yfinance location
  Sheet 2 — "Decision Maker Contacts" : Collapsible grouped contacts per company with clickable LinkedIn hyperlinks
"""

import os
import re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import Outline
import config

def export_consolidated_excel(companies: list, timestamp: str) -> str:
    """Exports 2-sheet consolidated Excel workbook: Verified Shortlist & Decision Maker Contacts."""
    os.makedirs(config.EXCEL_OUTPUT_DIR, exist_ok=True)
    filename = f"report_consolidated_{timestamp}.xlsx"
    filepath = os.path.join(config.EXCEL_OUTPUT_DIR, filename)

    wb = Workbook()
    
    # ── SHEET 1: Verified Shortlist ───────────────────────────
    ws1 = wb.active
    ws1.title = "Verified Shortlist"
    
    headers1 = [
        "Company Name", "Ticker", "Industry", "Revenue (TTM)", 
        "Net Income (TTM)", "Employee Count", "Location", "Business Summary", 
        "Website", "LinkedIn (Verified)", "LinkedIn Confirmed?"
    ]
    ws1.append(headers1)
    
    header_fill1 = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font1 = Font(color="FFFFFF", bold=True)
    
    for col in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.fill = header_fill1
        cell.font = header_font1
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for company in companies:
        is_confirmed = company.get("linkedin_confirmed")
        confirmed_str = "Verified" if is_confirmed else "Non-Verified"
        
        row = [
            company.get("name", "") or company.get("company_name", ""),
            company.get("ticker", ""),
            company.get("industry", "") or company.get("field", ""),
            company.get("revenue_ttm", "") or company.get("revenue", ""),
            company.get("net_income_ttm", "") or company.get("profit", ""),
            company.get("employee_count", "") or company.get("num_employees", ""),
            company.get("location", ""),
            company.get("business_summary", "") or company.get("summary", ""),
            company.get("website", "") or company.get("company_website", ""),
            company.get("linkedin_verified", "") or company.get("company_linkedin", ""),
            confirmed_str
        ]
        ws1.append(row)
        curr_row = ws1.max_row
        
        website = company.get("website", "") or company.get("company_website", "")
        if website and str(website).startswith("http"):
            ws1.cell(row=curr_row, column=9).value = f'=HYPERLINK("{website}", "{website}")'
            ws1.cell(row=curr_row, column=9).font = Font(color="0563C1", underline="single")
            
        linkedin_ver = company.get("linkedin_verified", "") or company.get("company_linkedin", "")
        if linkedin_ver and str(linkedin_ver).startswith("http"):
            ws1.cell(row=curr_row, column=10).value = f'=HYPERLINK("{linkedin_ver}", "{linkedin_ver}")'
            ws1.cell(row=curr_row, column=10).font = Font(color="0563C1", underline="single")

        conf_cell = ws1.cell(row=curr_row, column=11)
        if is_confirmed:
            conf_cell.font = Font(color="0F5132", bold=True)
            conf_cell.fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
        else:
            conf_cell.font = Font(color="842029", bold=True)
            conf_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        conf_cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws1.columns:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and not val_str.startswith("="):
                max_len = len(val_str)
        ws1.column_dimensions[col[0].column_letter].width = min(max(max_len + 3, 12), 60)

    # ── SHEET 2: Decision Maker Contacts (Collapsible Grouped Layout) ──
    ws2 = wb.create_sheet("Decision Maker Contacts")
    ws2.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)

    HEADERS2 = ["#", "Name", "Position", "LinkedIn (Person)", "Company LinkedIn", "Location"]
    widths2 = {"A": 6, "B": 28, "C": 50, "D": 45, "E": 45, "F": 18}
    for col_let, w in widths2.items():
        ws2.column_dimensions[col_let].width = w

    FNT = "Arial"
    company_font = Font(name=FNT, bold=True, size=12, color="FFFFFF")
    company_fill = PatternFill("solid", fgColor="2F5496")
    sub_hdr_font = Font(name=FNT, bold=True, size=10, color="2F5496")
    sub_hdr_fill = PatternFill("solid", fgColor="D6E4F0")
    data_font    = Font(name=FNT, size=10)
    link_font    = Font(name=FNT, size=10, color="0563C1", underline="single")
    alt_fill     = PatternFill("solid", fgColor="F2F2F2")
    border       = Border(
        left=Side("thin", color="B4C6E7"), right=Side("thin", color="B4C6E7"),
        top=Side("thin", color="B4C6E7"),  bottom=Side("thin", color="B4C6E7"),
    )
    wrap = Alignment(vertical="center", wrap_text=True)
    ctr  = Alignment(horizontal="center", vertical="center")

    row = 1
    for company_item in companies:
        comp_name = company_item.get("name", "") or company_item.get("company_name", "") or "—"
        co_url = company_item.get("linkedin_verified") or company_item.get("company_linkedin") or None
        if co_url and not str(co_url).startswith("http"):
            co_url = None
        co_label = comp_name if co_url else "—"

        comp_loc = company_item.get("location", "") or "—"
        dms = company_item.get("decision_makers", [])

        contacts = []
        if dms:
            for dm in dms:
                dname = dm.get("name") or "—"
                pos = dm.get("position") or dm.get("title") or "—"
                dm_url = dm.get("linkedin_url") or dm.get("profile_url") or None
                if dm_url and not str(dm_url).startswith("http"):
                    dm_url = None
                dm_label = dname if dm_url else "—"
                contacts.append({
                    "dm_name": dname,
                    "position": pos,
                    "dm_url": dm_url,
                    "dm_label": dm_label,
                })
        else:
            contacts.append({
                "dm_name": "—",
                "position": "—",
                "dm_url": None,
                "dm_label": "—",
            })

        # ── COMPANY HEADER ──────────────────────────────────────
        ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        hdr_cell = ws2.cell(row=row, column=1, value=f"  ▸ {comp_name}  ({len(contacts)} contacts)")
        hdr_cell.font, hdr_cell.fill, hdr_cell.alignment = company_font, company_fill, ctr
        for c in range(1, 7):
            ws2.cell(row=row, column=c).fill = company_fill
            ws2.cell(row=row, column=c).border = border
        ws2.row_dimensions[row].height = 28
        row += 1

        # ── SUB-HEADER ──────────────────────────────────────────
        first_data_row = row
        for ci, h in enumerate(HEADERS2, 1):
            cell = ws2.cell(row=row, column=ci, value=h)
            cell.font, cell.fill, cell.alignment, cell.border = sub_hdr_font, sub_hdr_fill, ctr, border
        row += 1

        # ── CONTACT ROWS ────────────────────────────────────────
        for idx, ct in enumerate(contacts, 1):
            # Col A: #
            c = ws2.cell(row=row, column=1, value=idx)
            c.font, c.border, c.alignment = data_font, border, ctr

            # Col B: Name
            c = ws2.cell(row=row, column=2, value=ct["dm_name"])
            c.font, c.border, c.alignment = data_font, border, wrap

            # Col C: Position
            c = ws2.cell(row=row, column=3, value=ct["position"])
            c.font, c.border, c.alignment = data_font, border, wrap

            # Col D: Person LinkedIn (clickable)
            if ct["dm_url"]:
                c = ws2.cell(row=row, column=4, value=ct["dm_label"])
                c.hyperlink = ct["dm_url"]
                c.font, c.border, c.alignment = link_font, border, wrap
            else:
                c = ws2.cell(row=row, column=4, value="—")
                c.font, c.border, c.alignment = data_font, border, wrap

            # Col E: Company LinkedIn (clickable)
            if co_url:
                c = ws2.cell(row=row, column=5, value=co_label)
                c.hyperlink = co_url
                c.font, c.border, c.alignment = link_font, border, wrap
            else:
                c = ws2.cell(row=row, column=5, value="—")
                c.font, c.border, c.alignment = data_font, border, wrap

            # Col F: Location
            c = ws2.cell(row=row, column=6, value=comp_loc)
            c.font, c.border, c.alignment = data_font, border, ctr

            # Alternating fill
            if idx % 2 == 0:
                for ci in range(1, 7):
                    ws2.cell(row=row, column=ci).fill = alt_fill

            row += 1

        last_data_row = row - 1

        # ── GROUP (collapsible) ─────────────────────────────────
        if last_data_row >= first_data_row:
            ws2.row_dimensions.group(first_data_row, last_data_row, hidden=True)
        row += 1  # spacer

    ws2.sheet_view.showGridLines = False

    wb.save(filepath)
    return filepath