import os
import sys
import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from output.excel_exporter import export_excel

def test_export_excel_verified_and_non_verified_status(tmp_path):
    sample_companies = [
        {
            "name": "Verified Tech Ltd",
            "ticker": "VT",
            "industry": "Software",
            "revenue_ttm": 1000000,
            "net_income_ttm": 200000,
            "employee_count": 500,
            "business_summary": "Tech company",
            "website": "https://verifiedtech.com",
            "linkedin_guessed": "https://linkedin.com/company/verified-tech",
            "linkedin_verified": "https://linkedin.com/company/verified-tech",
            "linkedin_confirmed": True
        },
        {
            "name": "Fallback Tech Ltd",
            "ticker": "FT",
            "industry": "Software",
            "revenue_ttm": 500000,
            "net_income_ttm": 100000,
            "employee_count": 200,
            "business_summary": "Fallback company",
            "website": "https://fallbacktech.com",
            "linkedin_guessed": "https://linkedin.com/company/fallback-tech",
            "linkedin_verified": "https://linkedin.com/company/fallback-tech",
            "linkedin_confirmed": False  # Fallback case
        }
    ]
    
    filepath = export_excel(sample_companies, timestamp="test_export")
    assert os.path.exists(filepath)
    
    wb = load_workbook(filepath)
    ws = wb.active
    
    # Header row is row 1. Row 2 is Verified Tech, Row 3 is Fallback Tech.
    # Column 10 is 'LinkedIn Confirmed?'
    row2_status = ws.cell(row=2, column=10).value
    row3_status = ws.cell(row=3, column=10).value
    
    assert row2_status == "Verified"
    assert row3_status == "Non-Verified"
